"""Synthetic DSD table workbook checks; no Excel or desktop input automation."""
from copy import deepcopy
from decimal import Decimal
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import load_workbook

from io_xlsx.package_guard import check_package
from xbrl_filing.disclosure_compare.dsd import read_dsd
from xbrl_filing.dsd_excel import (
    SCHEMA, META_SHEET, CELLS_SHEET, CELL_HEADERS, build_export_plan,
    cell_metadata_record, export_cell_value, export_dsd_excel, source_digest,
)
from xbrl_filing.dsd_excel import service


XML = '''<DOCUMENT><P>재무상태표 (단위: 천원)</P><TABLE><THEAD>
<TR><TH ROWSPAN="2">과목</TH><TH COLSPAN="2">2025년</TH></TR>
<TR><TH>기말</TH><TH>전기말</TH></TR></THEAD>
<TR><TD>국민은행</TD><TD>1,234</TD><TD>(20)</TD></TR>
<TR><TD>매출원가</TD><TD>0</TD><TD>-</TD></TR>
<TR><TD>이익률</TD><TD>12.50%</TD><TD></TD></TR>
<TR><TD>=HYPERLINK("https://invalid.example")</TD><TD>1234567890123456</TD><TD>00101</TD></TR>
</TABLE></DOCUMENT>'''


class DsdExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.source = self.root / '회사_2025.xml'
        self.source.write_text(XML, encoding='utf-8')
        self.target = self.root / '표.xlsx'

    def tearDown(self):
        self.temp.cleanup()

    def export(self):
        return export_dsd_excel(self.source, self.target)

    def test_all_tables_grid_merges_and_source_bound_metadata_survive(self):
        original = self.source.read_bytes()
        report = self.export()
        inventory = read_dsd(self.source)
        plan = build_export_plan(inventory)
        workbook = load_workbook(self.target, data_only=False)
        self.addCleanup(workbook.close)
        self.assertEqual(self.source.read_bytes(), original)
        self.assertEqual(report['sourceSha256'], source_digest(self.source))
        self.assertTrue(report['sourceUnchanged'])
        self.assertEqual(report['counts']['cells'], len(inventory['cells']))
        self.assertEqual(report['counts']['narratives'], 1)
        self.assertEqual(workbook['본문']['D3'].value, '재무상태표 (단위: 천원)')
        sheet = workbook[plan[0]['sheet']]
        self.assertEqual(str(sheet.merged_cells), 'A1:A2 B1:C1')
        self.assertEqual(sheet['B3'].value, 1234)
        self.assertEqual(sheet['C3'].value, -20)
        self.assertEqual(sheet['B4'].value, 0)
        self.assertEqual(sheet['C4'].value, '-')
        self.assertEqual(sheet['B5'].value, 0.125)
        self.assertIn('0.00%', sheet['B5'].number_format)
        self.assertIsNone(sheet['C5'].value)
        self.assertEqual(workbook[META_SHEET].sheet_state, 'veryHidden')
        self.assertEqual(workbook[CELLS_SHEET].sheet_state, 'veryHidden')
        metadata = list(workbook[META_SHEET].values)
        self.assertEqual(metadata[1], ('schema', SCHEMA))
        self.assertIn(('sourceSha256', report['sourceSha256']), metadata)
        self.assertIn(('sourceName', self.source.name), metadata)
        table = next(value for key, value in metadata if key == 'table')
        self.assertEqual(json.loads(table), plan[0])
        rows = list(workbook[CELLS_SHEET].values)
        self.assertEqual(rows[0], CELL_HEADERS)
        normalize = lambda values: [value if value is not None else '' for value in values]
        expected = [cell_metadata_record(cell, plan[0]['sheet']) for cell in inventory['cells']]
        self.assertEqual([normalize(row) for row in rows[1:]], [normalize(row) for row in expected])
        self.assertEqual(report['outputSha256'], source_digest(self.target))
        check_package(self.target)
        json.dumps(report, ensure_ascii=False)

    def test_formula_like_text_precision_and_leading_zeroes_are_not_changed(self):
        report = self.export()
        workbook = load_workbook(self.target, data_only=False)
        self.addCleanup(workbook.close)
        sheet = workbook[report['tables'][0]['sheet']]
        self.assertTrue(sheet['A6'].value.startswith('=HYPERLINK'))
        self.assertEqual(sheet['A6'].data_type, 's')
        self.assertIsNone(sheet['A6'].hyperlink)
        self.assertEqual(sheet['B6'].value, '1234567890123456')
        self.assertEqual(sheet['C6'].value, '00101')
        self.assertEqual(report['counts']['precisionTextCells'], 1)
        with ZipFile(self.target) as archive:
            data = b''.join(archive.read(name) for name in archive.namelist() if name.startswith('xl/worksheets/'))
            self.assertNotIn(b'<f>', data)
            self.assertFalse(any('externalLinks' in name for name in archive.namelist()))

    def test_header_year_remains_text_and_numeric_semantics_are_exact(self):
        for raw, number, expected_kind, expected_value in [
            ('2025', '2025', 'text', '2025'),
            ('(1.25%)', None, 'text', '(1.25%)'),
            ('-1.25%', '-0.0125', 'number', -0.0125),
            ('0.00000000000000000000001', '0.00000000000000000000001', 'number', 1e-23),
            ('0.1234567890123456', '0.1234567890123456', 'precision-text', '0.1234567890123456'),
        ]:
            with self.subTest(raw=raw):
                cell = {'text': raw, 'number': number, 'isHeader': raw == '2025'}
                value, kind, _ = export_cell_value(cell)
                self.assertEqual(kind, expected_kind)
                self.assertEqual(value, expected_value)
                if kind == 'number':
                    self.assertEqual(Decimal(str(value)), Decimal(number))

    def test_small_nonzero_numeric_values_do_not_display_as_zero(self):
        value, kind, number_format = export_cell_value({
            'text': '0.00000000000000000000001',
            'number': '0.00000000000000000000001', 'isHeader': False})
        self.assertEqual(value, 1e-23)
        self.assertEqual(kind, 'number')
        self.assertIn('E+00', number_format)

    def test_narratives_include_headings_inline_unknown_leaves_and_locations_without_duplicate_table_cells(self):
        self.source.write_text('<DOCUMENT><SECTION-1><TITLE>4. 재고자산</TITLE>'
            '<P>제품은 <SPAN>원가와 순실현가능가치</SPAN> 중 낮은 금액으로 평가합니다.</P>'
            '<UNRECOGNIZED>=DANGEROUS()</UNRECOGNIZED>'
            '<TABLE><TR><TH>품목</TH><TH>금액</TH></TR><TR><TD>원재료</TD><TD>100</TD></TR></TABLE>'
            '<P>평가손실은 비용에 포함합니다.</P></SECTION-1></DOCUMENT>', encoding='utf-8')
        report = self.export()
        book = load_workbook(self.target, data_only=False)
        self.addCleanup(book.close)
        sheet = book['본문']
        rows = list(sheet.iter_rows(min_row=3, values_only=True))
        self.assertEqual(report['counts']['narratives'], 4)
        self.assertEqual([row[3] for row in rows], [
            '4. 재고자산', '제품은 원가와 순실현가능가치 중 낮은 금액으로 평가합니다.',
            '=DANGEROUS()', '평가손실은 비용에 포함합니다.'])
        self.assertTrue(all(row[1] == self.source.name for row in rows))
        self.assertTrue(all(row[2] == '4. 재고자산' for row in rows))
        self.assertTrue(all('#p' in row[4] and 'XML' in row[4] for row in rows))
        self.assertEqual(sheet['D5'].data_type, 's')
        self.assertEqual(str(sheet.merged_cells), 'A1:E1')
        check_package(self.target)

    def test_long_narrative_is_rejected_instead_of_partially_exported(self):
        self.source.write_text('<DOCUMENT><P>' + '긴서술' * 11000 + '</P>'
            '<TABLE><TR><TD>1</TD></TR></TABLE></DOCUMENT>', encoding='utf-8')
        with self.assertRaisesRegex(ValueError, '32,767'):
            self.export()
        self.assertFalse(self.target.exists())

    def test_zip_members_duplicate_titles_and_unsafe_sheet_characters_are_preserved_as_provenance(self):
        source = self.root / '공시.dsd'
        with ZipFile(source, 'w', ZIP_DEFLATED) as archive:
            archive.writestr('folder1/a[very-long-source-name-for-sheet].xml', XML)
            archive.writestr('folder2/a[very-long-source-name-for-sheet].xml', XML)
        report = export_dsd_excel(source, self.target)
        titles = [table['sheet'] for table in report['tables']]
        self.assertEqual(len(titles), 2)
        self.assertNotEqual(titles[0].casefold(), titles[1].casefold())
        self.assertTrue(all(len(title) <= 31 and '[' not in title and ']' not in title for title in titles))
        self.assertEqual(report['counts']['cells'], 2 * len(read_dsd(self.source)['cells']))
        self.assertEqual([t['member'] for t in report['tables']],
                         ['folder1/a[very-long-source-name-for-sheet].xml', 'folder2/a[very-long-source-name-for-sheet].xml'])

    def test_existing_output_and_source_are_never_replaced(self):
        self.target.write_bytes(b'existing-output')
        with self.assertRaises(FileExistsError):
            self.export()
        self.assertEqual(self.target.read_bytes(), b'existing-output')
        with self.assertRaises(ValueError):
            export_dsd_excel(self.source, self.source)
        self.assertEqual(self.source.read_text(encoding='utf-8'), XML)

    def test_corrupt_xml_and_unsupported_binary_leave_no_workbook(self):
        for data in (b'not XML', b'<DOCUMENT><TABLE>', b'<!DOCTYPE a [<!ENTITY x "bad">]><a>&x;</a>'):
            with self.subTest(data=data):
                self.source.write_bytes(data)
                with self.assertRaises(ValueError):
                    self.export()
                self.assertFalse(self.target.exists())

    def test_long_text_and_wide_table_are_rejected_without_truncation(self):
        for table in ('<TR><TD>' + '가' * 32768 + '</TD></TR>',
                      '<TR>' + '<TD>1</TD>' * 81 + '</TR>',
                      '<TR><TD>' + '😀' * 16384 + '</TD></TR>'):
            with self.subTest(length=len(table)):
                self.source.write_text('<DOCUMENT><TABLE>' + table + '</TABLE></DOCUMENT>', encoding='utf-8')
                with self.assertRaises(ValueError):
                    self.export()
                self.assertFalse(self.target.exists())

    def test_too_many_tables_rejected_instead_of_silently_omitted(self):
        self.source.write_text('<DOCUMENT>' + '<TABLE><TR><TD>1</TD></TR></TABLE>' * 257 + '</DOCUMENT>', encoding='utf-8')
        with self.assertRaisesRegex(ValueError, '256'):
            self.export()
        self.assertFalse(self.target.exists())

    def test_changed_source_before_save_is_detected(self):
        real_digest = service.source_digest
        calls = 0
        def mutate(path):
            nonlocal calls
            calls += 1
            if calls == 2:
                self.source.write_text(XML.replace('1,234', '9,999'), encoding='utf-8')
            return real_digest(path)
        with patch.object(service, 'source_digest', side_effect=mutate):
            with self.assertRaisesRegex(ValueError, '바뀌'):
                self.export()
        self.assertFalse(self.target.exists())

    def test_changed_source_during_save_and_save_errors_clean_temporary_files(self):
        real_save = service.Workbook.save
        def mutate(workbook, filename):
            real_save(workbook, filename)
            self.source.write_text(XML.replace('1,234', '9,999'), encoding='utf-8')
        with patch.object(service.Workbook, 'save', new=mutate):
            with self.assertRaisesRegex(ValueError, '바뀌'):
                self.export()
        self.assertFalse(self.target.exists())
        self.assertFalse(list(self.root.glob('.dsd-excel-*')))
        self.source.write_text(XML, encoding='utf-8')
        with patch.object(service.Workbook, 'save', side_effect=OSError('synthetic save failure')):
            with self.assertRaises(OSError):
                self.export()
        self.assertFalse(self.target.exists())
        self.assertFalse(list(self.root.glob('.dsd-excel-*')))

    def test_output_race_does_not_replace_newly_created_file(self):
        real_commit = service._commit
        def race(temp, target):
            target.write_bytes(b'other-writer')
            real_commit(temp, target)
        with patch.object(service, '_commit', side_effect=race):
            with self.assertRaises(FileExistsError):
                self.export()
        self.assertEqual(self.target.read_bytes(), b'other-writer')
        self.assertFalse(list(self.root.glob('.dsd-excel-*')))

    def test_invalid_coverage_and_global_merged_grid_are_rejected(self):
        inventory = read_dsd(self.source)
        with self.subTest('uncovered'):
            broken = deepcopy(inventory)
            broken['tables'][0]['cellIds'].pop()
            with self.assertRaises(ValueError):
                build_export_plan(broken)
        with self.subTest('duplicate'):
            broken = deepcopy(inventory)
            broken['tables'][0]['cellIds'].append(broken['tables'][0]['cellIds'][0])
            with self.assertRaises(ValueError):
                build_export_plan(broken)
        with self.subTest('merged memory budget'):
            broken = deepcopy(inventory)
            for cell in broken['cells']:
                cell['source']['rowspan'] = 100000
                cell['source']['colspan'] = 80
            with self.assertRaisesRegex(ValueError, '1,000,000'):
                build_export_plan(broken)


if __name__ == '__main__':
    unittest.main()

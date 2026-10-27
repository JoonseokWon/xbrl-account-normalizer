import csv
from datetime import date
import hashlib
from io import StringIO
import json
from pathlib import Path
import sys
import tempfile
import unittest
import warnings
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from core.errors import InputError
from xbrl_filing.inputs import HEADERS,create_template,read_input,read_work,write_work


class FilingInputsTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()
    def row(self,**changes):
        data=dict(zip(HEADERS,['0001','BS','0010','현금및현금성자산','ifrs:Cash','(1,234.50)','KRW','1000','-3','instant','2026-12-31','','','{}']))
        return data|changes
    def csv(self,rows=None,name='input.csv',encoding='utf-8-sig'):
        stream=StringIO(newline='');writer=csv.writer(stream);writer.writerow(HEADERS)
        for row in (rows if rows is not None else [self.row()]):writer.writerow([row.get(key,'') for key in HEADERS])
        path=self.root/name;path.write_bytes(stream.getvalue().encode(encoding));return path
    def report(self):
        report=read_input(self.csv());report.update(taxonomyFingerprint='abc123',taxonomyEntryPoint='official.xsd')
        report['facts'][0].update(approved=True,reviewer='확인자');return report

    def test_blank_template_does_not_import_the_separate_examples(self):
        from openpyxl import load_workbook
        path=create_template(self.root/'template.xlsx');report=read_input(path)
        self.assertEqual(report['facts'],[]);self.assertEqual(report['entity']['identifier'],'')
        workbook=load_workbook(path)
        self.assertEqual(workbook['공시항목'].freeze_panes,'A2')
        self.assertEqual(workbook['공시항목'].max_row,1)
        self.assertGreater(workbook['작성예시'].max_row,1)
        self.assertEqual(workbook['작성예시']['A2'].value,'0001');workbook.close()
        old=path.read_bytes()
        with self.assertRaises(FileExistsError):create_template(path)
        self.assertEqual(path.read_bytes(),old)

    def test_excel_dates_identifiers_and_display_scale_are_preserved(self):
        from openpyxl import load_workbook
        path=create_template(self.root/'filled.xlsx');workbook=load_workbook(path)
        values={'회사명':'입력 회사','기업식별자':'00123456','식별자체계URI':'urn:company:id','당기시작일':date(2026,1,1),'당기종료일':date(2026,12,31),'연결구분':'연결','비교기간JSON':'[{"periodStart":"2025-01-01","periodEnd":"2025-12-31"}]'}
        for row in workbook['공시정보'].iter_rows(min_row=2):row[1].value=values[row[0].value]
        record=self.row(기준일=date(2026,12,31));workbook['공시항목'].append([record[key] for key in HEADERS]);workbook.save(path);workbook.close()
        report=read_input(path);fact=report['facts'][0]
        self.assertEqual(report['entity']['identifier'],'00123456');self.assertEqual(report['reporting']['consolidation'],'연결')
        self.assertEqual(fact['id'],'f000001');self.assertEqual(fact['originalId'],'0001');self.assertEqual(fact['originalCode'],'0010')
        self.assertEqual(fact['value'],'-1234500');self.assertEqual(fact['decimals'],'-3');self.assertEqual(fact['displayScale'],1000)
        self.assertEqual(fact['originalValue'],'(1,234.50)');self.assertEqual(fact['period']['instant'],'2026-12-31')
        self.assertEqual(fact['source'],'filled.xlsx:공시항목:2');self.assertFalse(fact['approved'])
        self.assertEqual(report['inputs'][0]['sha256'],hashlib.sha256(path.read_bytes()).hexdigest())

    def test_numeric_excel_codes_are_rejected_instead_of_inventing_leading_zeroes(self):
        from openpyxl import load_workbook
        for field in ('항목ID','원본코드'):
            path=create_template(self.root/(field+'.xlsx'));workbook=load_workbook(path)
            row=self.row(**{field:1});workbook['공시항목'].append([row[key] for key in HEADERS]);workbook.save(path);workbook.close()
            with self.subTest(field=field),self.assertRaises(InputError):read_input(path)

    def test_csv_has_explicit_blank_company_scope_and_cp949_support(self):
        report=read_input(self.csv(encoding='cp949'))
        self.assertEqual(report['entity']['name'],'');self.assertEqual(report['reporting']['periodStart'],'')
        self.assertEqual(report['facts'][0]['originalCode'],'0010');self.assertEqual(report['facts'][0]['source'],'input.csv:CSV:2')

    def test_scale_and_accuracy_are_independent(self):
        report=read_input(self.csv([self.row(값='1,234.50',정확도='0'),self.row(값='1,234.50',정확도='-3')]))
        self.assertEqual([fact['value'] for fact in report['facts']],['1234500','1234500'])
        self.assertEqual([fact['decimals'] for fact in report['facts']],['0','-3'])

    def test_invalid_numbers_dates_units_scales_and_dimensions_are_rejected(self):
        invalid=[{'값':'1,23'},{'값':'NaN'},{'값':'1e10'},{'값':True},{'표시배율':'100'},{'기준일':'20261231'},
                 {'단위':'USD'},{'차원JSON':'[]'},{'차원JSON':'{"axis":"member","axis":"other"}'},{'정확도':'999'}]
        for index,changes in enumerate(invalid):
            path=self.csv([self.row(**changes)],str(index)+'.csv')
            with self.subTest(changes=changes),self.assertRaises(InputError):read_input(path)

    def test_text_and_duration_facts_do_not_manufacture_numeric_values(self):
        row=self.row(재무제표='NOTE',값='확정한 주석 설명',단위='',표시배율='1',정확도='',기간형식='duration',기준일='',시작일='2026-01-01',종료일='2026-12-31')
        report=read_input(self.csv([row]));fact=report['facts'][0]
        self.assertEqual(fact['value'],'확정한 주석 설명');self.assertEqual(fact['unit'],'');self.assertEqual(fact['period']['kind'],'duration')

    def test_generic_json_revokes_approvals_without_rescaling_a_saved_value(self):
        report=self.report();path=self.root/'input.json';path.write_text(json.dumps(report,ensure_ascii=False),encoding='utf-8')
        imported=read_input(path)
        self.assertFalse(imported['facts'][0]['approved']);self.assertEqual(imported['facts'][0]['reviewer'],'')
        self.assertEqual(imported['facts'][0]['value'],'-1234500');self.assertEqual(imported['facts'][0]['displayScale'],1000)

    def test_work_preserves_unsigned_reviews_and_original_source_metadata(self):
        report=self.report();path=write_work(report,self.root/'review.acfiling');loaded=read_work(path)
        self.assertEqual(loaded,report);self.assertTrue(loaded['facts'][0]['approved'])
        self.assertEqual(loaded['taxonomyFingerprint'],'abc123')
        report['facts'][0]['reviewer']='다른 확인자';write_work(report,path)
        self.assertEqual(read_work(path)['facts'][0]['reviewer'],'다른 확인자')

    def test_failed_atomic_work_save_preserves_previous_file_and_removes_temp(self):
        report=self.report();path=write_work(report,self.root/'review.acfiling');before=path.read_bytes()
        with patch('xbrl_filing.inputs.os.replace',side_effect=OSError('busy')):
            with self.assertRaises(OSError):write_work(report,path)
        self.assertEqual(path.read_bytes(),before);self.assertEqual(list(self.root.glob('.filing-work-*')),[])

    def test_duplicate_json_keys_nonfinite_numbers_and_wrong_shapes_are_rejected(self):
        for index,text in enumerate(('{"schema":"a","schema":"b"}','{"value":NaN}','{"value":Infinity}','{"value":1e999}','[]')):
            path=self.root/(str(index)+'.json');path.write_text(text,encoding='utf-8')
            with self.subTest(text=text),self.assertRaises(InputError):read_input(path)

    def test_formulas_in_unused_example_sheet_are_also_rejected(self):
        from openpyxl import load_workbook
        path=create_template(self.root/'formula.xlsx');workbook=load_workbook(path);workbook['작성예시']['F2']='=1+2';workbook.save(path);workbook.close()
        with self.assertRaises(InputError):read_input(path)

    def test_csv_formula_strings_and_oversized_xlsx_text_are_rejected(self):
        path=self.csv([self.row(단위='',표시배율='1',값='=SUM(A1:A2)')])
        with self.assertRaises(InputError):read_input(path)
        path=self.root/'text.xlsx'
        with ZipFile(path,'w') as archive:archive.writestr('xl/worksheets/sheet1.xml','<worksheet><row r="1"><c r="A1"><t>text-too-long</t></c></row></worksheet>')
        with patch('xbrl_filing.inputs.MAX_TEXT',5),self.assertRaises(InputError):read_input(path)

    def test_unsafe_xlsx_archives_are_rejected_before_workbook_loading(self):
        cases=[ [('xl/a.xml',b'<root/>'),('xl/a.xml',b'<root/>')],
                [('../escape.xml',b'<root/>')],
                [('xl/vbaProject.bin',b'bad')],
                [('xl/_rels/workbook.xml.rels',b'<Relationships><Relationship TargetMode="External" Target="https://example.com"/></Relationships>')],
                [('xl/workbook.xml','<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE a [<!ENTITY x "bad">]><a>&x;</a>'.encode('utf-16'))] ]
        for index,entries in enumerate(cases):
            path=self.root/('unsafe'+str(index)+'.xlsx')
            with warnings.catch_warnings():
                warnings.simplefilter('ignore',UserWarning)
                with ZipFile(path,'w') as archive:
                    for name,data in entries:archive.writestr(name,data)
            with self.subTest(index=index),self.assertRaises(InputError):read_input(path)

    def test_input_and_fact_limits_are_enforced(self):
        path=self.csv([self.row(),self.row()])
        with patch('xbrl_filing.inputs.MAX_FACTS',1),self.assertRaises(InputError):read_input(path)
        with patch('xbrl_filing.inputs.MAX_BYTES',20),self.assertRaises(InputError):read_input(path)


if __name__=='__main__':unittest.main()

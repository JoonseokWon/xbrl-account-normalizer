from copy import deepcopy
import hashlib
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook

from xbrl_filing.disclosure_compare.dsd import read_dsd
from xbrl_filing.dsd_roundtrip import build_workbook_mapping, export_workbook_to_dsd, inspect_workbook_cells


class DsdRoundtripTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.template = self.root/'original.dsd'
        self.xml = b'<?xml version="1.0" encoding="UTF-8"?><DOCUMENT flag="keep"><P>keep text</P><TABLE><TR><TH>item</TH><TH>amount</TH></TR><TR><TD>cash</TD><TD>(123)</TD></TR><TR><TD>rate</TD><TD>5%</TD></TR><TR><TD>blank</TD><TD/></TR></TABLE><UNKNOWN attr="keep"/></DOCUMENT>'
        with ZipFile(self.template, 'w', ZIP_DEFLATED) as zipped:
            zipped.writestr('Contents.XML', self.xml)
            zipped.writestr('meta.xml', '<META stable="yes"/>')
            zipped.writestr('assets/image.png', b'not executed')
            zipped.comment = b'keep archive comment'
        self.workbook = self.root/'source.xlsx'
        book = Workbook(); book.active.title = '정산표'
        book.active['A1'] = 200
        book.active['B1'] = 0.25; book.active['B1'].number_format = '0.0%'
        book.active['C1'] = '=literal'; book.active['C1'].data_type = 's'
        book.save(self.workbook); book.close()
        self.inventory = read_dsd(self.template)
        self.cash = next(c for c in self.inventory['cells'] if c['text'] == '(123)')['id']
        self.rate = next(c for c in self.inventory['cells'] if c['text'] == '5%')['id']

    def tearDown(self):
        self.temp.cleanup()

    def mapping(self, identifier=None, address='A1', scale='1'):
        return build_workbook_mapping(self.template, self.workbook, [{'sheet': '정산표', 'cell': address,
            'dsdCellId': identifier or self.cash, 'confirmed': True, 'scale': scale}])

    def patch_workbook_xml(self, before, after):
        with ZipFile(self.workbook) as source:
            entries = [(entry, source.read(entry)) for entry in source.infolist()]
        with ZipFile(self.workbook, 'w') as output:
            for entry, data in entries:
                if entry.filename == 'xl/worksheets/sheet1.xml':
                    self.assertIn(before, data)
                    data = data.replace(before, after)
                output.writestr(entry, data)

    def test_explicit_arbitrary_workbook_values_preserve_unknown_parts(self):
        before = self.template.read_bytes()
        result = export_workbook_to_dsd(self.template, self.workbook, self.root/'out.dsd', self.mapping())
        self.assertEqual(result['changedCellCount'], 1)
        self.assertTrue(result['reopenedAndChecked'])
        self.assertFalse(result['officialValidated'])
        with ZipFile(self.template) as a, ZipFile(self.root/'out.dsd') as b:
            self.assertEqual(a.namelist(), b.namelist())
            self.assertEqual(a.comment, b.comment)
            self.assertEqual(a.read('meta.xml'), b.read('meta.xml'))
            self.assertEqual(a.read('assets/image.png'), b.read('assets/image.png'))
            xml = b.read('Contents.XML')
            self.assertIn(b'<UNKNOWN attr="keep"/>', xml)
            self.assertIn(b'<P>keep text</P>', xml)
        self.assertEqual(self.template.read_bytes(), before)
        self.assertIn('200', [c['text'] for c in read_dsd(self.root/'out.dsd')['cells']])

    def test_scale_only_explicit_numeric_and_percent_display_preserved(self):
        export_workbook_to_dsd(self.template, self.workbook, self.root/'scale.dsd', self.mapping(scale='1000'))
        self.assertIn('200000', [c['text'] for c in read_dsd(self.root/'scale.dsd')['cells']])
        export_workbook_to_dsd(self.template, self.workbook, self.root/'percent.dsd', self.mapping(self.rate, 'B1'))
        self.assertIn('25%', [c['text'] for c in read_dsd(self.root/'percent.dsd')['cells']])
        with self.assertRaisesRegex(ValueError, '숫자'):
            export_workbook_to_dsd(self.template, self.workbook, self.root/'badscale.dsd', self.mapping(address='C1', scale='1000'))

    def test_formula_like_text_stays_escaped_data(self):
        export_workbook_to_dsd(self.template, self.workbook, self.root/'text.dsd', self.mapping(address='C1'))
        self.assertIn('=literal', [c['text'] for c in read_dsd(self.root/'text.dsd')['cells']])

    def test_blank_clearing_is_distinct_from_zero(self):
        export_workbook_to_dsd(self.template, self.workbook, self.root/'blank.dsd', self.mapping(address='D1'))
        target = next(c for c in read_dsd(self.root/'blank.dsd')['cells'] if c['source']['row'] == 2 and c['source']['column'] == 2)
        self.assertEqual(target['text'], '')
        self.assertIsNone(target['number'])

    def test_source_change_or_unbound_list_cannot_apply(self):
        mapping = self.mapping()
        book = load_workbook(self.workbook); book.active['A1'] = 999; book.save(self.workbook); book.close()
        with self.assertRaisesRegex(ValueError, '해시'):
            export_workbook_to_dsd(self.template, self.workbook, self.root/'changed.dsd', mapping)
        with self.assertRaises(ValueError):
            export_workbook_to_dsd(self.template, self.workbook, self.root/'list.dsd', mapping['mappings'])

    def test_duplicate_unconfirmed_and_merged_nonanchor_rejected(self):
        item = {'sheet':'정산표','cell':'A1','dsdCellId':self.cash,'confirmed':False}
        with self.assertRaises(ValueError): build_workbook_mapping(self.template, self.workbook, [item])
        item['confirmed'] = True
        with self.assertRaises(ValueError): build_workbook_mapping(self.template, self.workbook, [item,item])
        book = load_workbook(self.workbook); book.active.merge_cells('A1:A2'); book.save(self.workbook); book.close()
        item['cell'] = 'A2'
        with self.assertRaisesRegex(ValueError, '병합'):
            build_workbook_mapping(self.template, self.workbook, [item])

    def test_actual_formula_is_rejected_before_read(self):
        book = load_workbook(self.workbook); book.active['A1'] = '=1+1'; book.save(self.workbook); book.close()
        with self.assertRaises(ValueError): inspect_workbook_cells(self.workbook)

    def test_inspection_includes_source_hash_and_number_format(self):
        result = inspect_workbook_cells(self.workbook)
        self.assertEqual(result['sourceHash'], hashlib.sha256(self.workbook.read_bytes()).hexdigest())
        self.assertEqual(result['sheets'], ['정산표'])
        percent = next(c for c in result['cells'] if c['cell'] == 'B1')
        self.assertEqual(percent['numberFormat'], '0.0%')

    def test_raw_high_precision_decimal_is_not_rounded_by_float_reader(self):
        raw = b'1.000000000000000001'
        self.patch_workbook_xml(b'<v>200</v>', b'<v>' + raw + b'</v>')
        inspected = inspect_workbook_cells(self.workbook)
        actual = next(c for c in inspected['cells'] if c['cell'] == 'A1')
        self.assertEqual(actual['value'], raw.decode())
        export_workbook_to_dsd(self.template, self.workbook, self.root/'precise.dsd', self.mapping())
        self.assertIn(raw.decode(), [c['text'] for c in read_dsd(self.root/'precise.dsd')['cells']])

    def test_unbounded_numeric_exponent_and_duplicate_coordinates_rejected(self):
        self.patch_workbook_xml(b'<v>200</v>', b'<v>1e10000000</v>')
        with self.assertRaises(ValueError): inspect_workbook_cells(self.workbook)
        # Restore the extreme exponent and insert a second physical A1.
        self.patch_workbook_xml(b'<v>1e10000000</v>', b'<v>200</v>')
        self.patch_workbook_xml(b'</row>', b'<c r="a1" t="n"><v>200</v></c></row>')
        with self.assertRaisesRegex(ValueError, '중복'):
            inspect_workbook_cells(self.workbook)

    def test_literal_percent_format_does_not_multiply_value(self):
        book = load_workbook(self.workbook); book.active['B1'].number_format = '0.00"%"'; book.save(self.workbook); book.close()
        export_workbook_to_dsd(self.template, self.workbook, self.root/'literal.dsd', self.mapping(self.rate, 'B1'))
        self.assertIn('0.25', [c['text'] for c in read_dsd(self.root/'literal.dsd')['cells']])
        book = load_workbook(self.workbook); book.active['B1'].number_format = '[>=1]0.0%;0.0%'; book.save(self.workbook); book.close()
        with self.assertRaisesRegex(ValueError, '조건부'):
            export_workbook_to_dsd(self.template, self.workbook, self.root/'conditional.dsd', self.mapping(self.rate, 'B1'))

    def test_mixed_content_not_silently_destroyed(self):
        self.template.write_bytes(self.xml.replace(b'<TD>(123)</TD>', b'<TD><SPAN>(1</SPAN><B>23)</B></TD>'))
        self.inventory = read_dsd(self.template)
        self.cash = next(c for c in self.inventory['cells'] if c['text'] == '(123)')['id']
        with self.assertRaisesRegex(ValueError, '중첩'):
            export_workbook_to_dsd(self.template, self.workbook, self.root/'mixed.dsd', self.mapping())
        self.assertFalse((self.root/'mixed.dsd').exists())

    def test_single_text_formatting_chain_preserved(self):
        self.template.write_bytes(self.xml.replace(b'<TD>(123)</TD>', b'<TD><P><SPAN style="keep">(123)</SPAN></P></TD>'))
        self.cash = next(c for c in read_dsd(self.template)['cells'] if c['text'] == '(123)')['id']
        export_workbook_to_dsd(self.template, self.workbook, self.root/'styled.dsd', self.mapping())
        self.assertIn(b'<P><SPAN style="keep">200</SPAN></P>', (self.root/'styled.dsd').read_bytes())

    def test_default_xml_namespace_preserved(self):
        self.template.write_bytes(self.xml.replace(b'<DOCUMENT flag="keep">', b'<DOCUMENT xmlns="urn:test:dart" flag="keep">'))
        self.cash = next(c for c in read_dsd(self.template)['cells'] if c['text'] == '(123)')['id']
        export_workbook_to_dsd(self.template, self.workbook, self.root/'namespaced.dsd', self.mapping())
        self.assertIn(b'xmlns="urn:test:dart"', (self.root/'namespaced.dsd').read_bytes())
        self.assertIn('200', [c['text'] for c in read_dsd(self.root/'namespaced.dsd')['cells']])

    def test_unchanged_derived_workbook_preserves_original_bytes_then_edit_updates(self):
        from xbrl_filing.dsd_excel import export_dsd_excel
        extracted = self.root/'derived.xlsx'
        export_dsd_excel(self.template, extracted)
        export_workbook_to_dsd(self.template, extracted, self.root/'same.dsd')
        with ZipFile(self.template) as a, ZipFile(self.root/'same.dsd') as b:
            self.assertEqual(a.read('Contents.XML'), b.read('Contents.XML'))
        book = load_workbook(extracted)
        data = next(s for s in book if s.title not in ('_DSD_META','_DSD_CELLS','변환안내'))
        # A summary sheet may precede data; the cell metadata identifies it.
        for values in book['_DSD_CELLS'].iter_rows(min_row=2,values_only=True):
            if values[0] == self.cash:
                book[values[1]][values[2]] = 888
        book.save(extracted); book.close()
        result = export_workbook_to_dsd(self.template, extracted, self.root/'edited.dsd')
        self.assertEqual(result['changedCellCount'], 1)
        self.assertIn('888', [c['text'] for c in read_dsd(self.root/'edited.dsd')['cells']])

    def test_tampered_derived_provenance_rejected(self):
        from xbrl_filing.dsd_excel import export_dsd_excel
        extracted = self.root/'derived.xlsx'; export_dsd_excel(self.template, extracted)
        book = load_workbook(extracted); book['_DSD_CELLS']['A2'] = 'malicious'; book.save(extracted); book.close()
        with self.assertRaisesRegex(ValueError, '출처'):
            export_workbook_to_dsd(self.template, extracted, self.root/'tampered.dsd')

    def test_derived_merge_structure_change_rejected(self):
        from xbrl_filing.dsd_excel import export_dsd_excel
        extracted = self.root/'merged.xlsx'; export_dsd_excel(self.template, extracted)
        book = load_workbook(extracted)
        values = next(book['_DSD_CELLS'].iter_rows(min_row=2,values_only=True))
        book[values[1]].merge_cells('A1:B1'); book.save(extracted); book.close()
        with self.assertRaisesRegex(ValueError, '병합'):
            export_workbook_to_dsd(self.template, extracted, self.root/'merge-bad.dsd')

    def test_output_failure_keeps_sources_and_no_partial_target(self):
        before_template, before_book = self.template.read_bytes(), self.workbook.read_bytes()
        with patch('xbrl_filing.dsd_roundtrip.service.os.link', side_effect=OSError('simulated failure')):
            with self.assertRaises(OSError):
                export_workbook_to_dsd(self.template, self.workbook, self.root/'failed.dsd', self.mapping())
        self.assertFalse((self.root/'failed.dsd').exists())
        self.assertFalse(list(self.root.glob('.dsd-draft-*')))
        self.assertEqual(self.template.read_bytes(), before_template)
        self.assertEqual(self.workbook.read_bytes(), before_book)

    def test_source_changed_during_generation_blocks_commit(self):
        from xbrl_filing.dsd_roundtrip.service import _patch_xml
        mapping = self.mapping()
        def mutate(*args):
            result = _patch_xml(*args)
            self.workbook.write_bytes(b'changed during operation')
            return result
        with patch('xbrl_filing.dsd_roundtrip.service._patch_xml', side_effect=mutate):
            with self.assertRaisesRegex(ValueError, '변경'):
                export_workbook_to_dsd(self.template, self.workbook, self.root/'stale.dsd', mapping)
        self.assertFalse((self.root/'stale.dsd').exists())
        self.assertFalse(list(self.root.glob('.dsd-draft-*')))

    def test_utf16_xml_template_keeps_encoding_and_unknown_nodes(self):
        text = self.xml.decode('utf-8').replace('UTF-8', 'UTF-16')
        self.template.write_bytes(text.encode('utf-16'))
        self.cash = next(c for c in read_dsd(self.template)['cells'] if c['text'] == '(123)')['id']
        export_workbook_to_dsd(self.template, self.workbook, self.root/'utf16.dsd', self.mapping())
        data = (self.root/'utf16.dsd').read_bytes()
        self.assertTrue(data.startswith((b'\xff\xfe', b'\xfe\xff')))
        self.assertIn('<UNKNOWN attr="keep"/>', data.decode('utf-16'))
        self.assertIn('200', [c['text'] for c in read_dsd(self.root/'utf16.dsd')['cells']])

    def test_existing_target_and_extension_changes_rejected(self):
        target = self.root/'exists.dsd'; target.write_bytes(b'keep')
        with self.assertRaises(FileExistsError): export_workbook_to_dsd(self.template, self.workbook, target, self.mapping())
        self.assertEqual(target.read_bytes(), b'keep')
        with self.assertRaises(ValueError): export_workbook_to_dsd(self.template, self.workbook, self.root/'bad.xml', self.mapping())


if __name__ == '__main__':
    unittest.main()

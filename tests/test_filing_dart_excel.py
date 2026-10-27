import hashlib
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import warnings
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from openpyxl import Workbook
from core.errors import InputError
from xbrl_filing import dart_excel as reader


class DartExcelTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory();self.root = Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()
    def book(self, *, duplicate=False, formula=False, period='instant', abstract='', missing=False):
        wb = Workbook();ws = wb.active;ws.title = 'Concepts'
        ws.append(['#','prefix','name','id','type','enumerations','substitutionGroup','balance','periodType','abstract','nillable','typedDomainRef','ext. enum domain','ext. enum linkrole','ext. enum headUsable','comment'])
        row = [1,'ifrs-full','Assets','ifrs-full_Assets','xbrli:monetaryItemType','','xbrli:item','debit',period,abstract,'true']
        if formula:row[2] = '=1+1'
        ws.append(row)
        if duplicate:ws.append(row)
        ws.append([2,'ifrs-full','Revenue','ifrs-full_Revenue','xbrli:monetaryItemType','','xbrli:item','credit','duration','','true'])
        ws.append([3,'ifrs-full','Unused','ifrs-full_Unused','xbrli:stringItemType','','xbrli:item','','duration','','true'])
        if missing:ws['B1'] = 'wrong'
        labels = wb.create_sheet('Label Link')
        labels.append(['LinkRole','http://www.xbrl.org/2003/role/link']);labels.append(['Definition','']);labels.append([None,None,None,'ko'])
        labels.append(['#','prefix','name','label','deprecatedDateLabel','deprecatedLabel'])
        labels.append([1,'ifrs-full','Assets','자산','',''])
        labels.append([2,'ifrs-full','Unused','이전 항목 (사용중단)','2024-01-01',''])
        labels.append([3,'','','','',''])
        roles = wb.create_sheet('RoleTypes');roles.append(['#','id','roleURI','usedOn','definition','Defined Schema Systemid','comment'])
        roles.append([1,'role-D210005','http://example.invalid/role/D210005','link:presentationLink','[D210005] 재무상태표 - 별도','real-dart.xsd',''])
        listing = wb.create_sheet('Role 일람표');listing.append(['test metadata']);listing.append([None,'업종','구분','별도/연결 구분','Link','주석 명칭(Label)'])
        listing.append([None,'비금융업','재무제표 본문','별도',"Presentation Link'!A1",'[D210005] 재무상태표 - 별도'])
        pre = wb.create_sheet('Presentation Link');pre.append(['LinkRole','http://example.invalid/role/D210005']);pre.append(['Definition','[D210005] 재무상태표 - 별도']);pre.append(['prefix','name','label'])
        pre.append(['ifrs-full','Assets','자산']);pre.append(['ifrs-full','Assets','자산']);pre.append(['ifrs-full','Revenue','수익'])
        path = self.root/'taxonomy.xlsx';wb.save(path);wb.close();return path
    def rewrite(self, path, name, data):
        with ZipFile(path) as archive:entries = [(e.filename,archive.read(e)) for e in archive.infolist()]
        with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
            for key,value in entries:archive.writestr(key, data if key == name else value)
    def test_exact_metadata_and_no_invented_namespace(self):
        path = self.book();value = reader.read_dart_taxonomy(path)
        self.assertEqual(value['validationScope'], 'concept-metadata-only')
        self.assertEqual(value['namespaces'], {})
        self.assertTrue(value['namespaceResolutionRequired'])
        self.assertEqual(value['source']['name'], path.name)
        self.assertEqual(value['source']['sha256'], hashlib.sha256(path.read_bytes()).hexdigest())
        asset = value['concepts']['ifrs-full:Assets']
        self.assertEqual((asset['label'],asset['namespace'],asset['periodType'],asset['kind']), ('자산','','instant','monetary'))
        self.assertTrue(asset['numeric']);self.assertFalse(asset['abstract'])
        self.assertTrue(value['concepts']['ifrs-full:Unused']['deprecated'])
        role = value['roles']['http://example.invalid/role/D210005']
        self.assertEqual((role['industry'],role['scope'],role['category']), ('비금융업','별도','재무제표 본문'))
        self.assertEqual(role['concepts'], ['ifrs-full:Assets','ifrs-full:Revenue'])
    def test_wrong_headers(self):
        with self.assertRaises(InputError):reader.read_dart_taxonomy(self.book(missing=True))
    def test_identical_duplicate_deduplicated(self):
        value=reader.read_dart_taxonomy(self.book(duplicate=True))
        self.assertEqual(len(value['concepts']),3)
        self.assertEqual(value['excludedConcepts'],[])
    def test_conflicting_duplicate_excluded_without_guessing_aliases(self):
        path=self.book(duplicate=True)
        from openpyxl import load_workbook
        workbook=load_workbook(path);workbook['Concepts']['E3']='custom:monetaryItemType'
        workbook.save(path);workbook.close()
        value=reader.read_dart_taxonomy(path)
        self.assertNotIn('ifrs-full:Assets',value['concepts'])
        self.assertEqual(value['excludedConcepts'],[{'qname':'ifrs-full:Assets','reason':'conflicting-declaration'}])
    def test_formula_not_evaluated(self):
        with self.assertRaises(InputError):reader.read_dart_taxonomy(self.book(formula=True))
    def test_invalid_period_and_boolean(self):
        for change in ({'period':'random'},{'abstract':'maybe'}):
            with self.subTest(change=change), self.assertRaises(InputError):reader.read_dart_taxonomy(self.book(**change))
    def test_file_and_expansion_limits(self):
        path = self.book()
        with patch.object(reader,'MAX_BYTES',100),self.assertRaises(InputError):reader.read_dart_taxonomy(path)
        with patch.object(reader,'MAX_EXPANDED',100),self.assertRaises(InputError):reader.read_dart_taxonomy(path)
        with patch.object(reader,'MAX_ENTRY_BYTES',100),self.assertRaises(InputError):reader.read_dart_taxonomy(path)
    def test_concept_and_relationship_limits(self):
        path = self.book()
        with patch.object(reader,'MAX_CONCEPTS',1),self.assertRaises(InputError):reader.read_dart_taxonomy(path)
        with patch.object(reader,'MAX_RELATIONSHIPS',1),self.assertRaises(InputError):reader.read_dart_taxonomy(path)
    def test_dtd_rejected_including_utf16(self):
        for encoding in ('utf-8','utf-16'):
            path = self.book();payload = ('<?xml version="1.0" encoding="'+encoding+'"?><!DOCTYPE root [<!ENTITY x "bad">]><root>&x;</root>').encode(encoding)
            self.rewrite(path,'xl/worksheets/sheet1.xml',payload)
            with self.subTest(encoding=encoding),self.assertRaises(InputError):reader.read_dart_taxonomy(path)
    def test_zip_path_duplicate_and_macro_rejected(self):
        for name in ('../bad.xml','C:/bad.xml','bad\\bad.xml','xl/vbaProject.bin','xl/workbook.xml','XL/workbook.xml'):
            path = self.book()
            with warnings.catch_warnings():
                warnings.simplefilter('ignore')
                entry=ZipInfo('placeholder');entry.filename=name
                with ZipFile(path,'a') as archive:archive.writestr(entry,b'<bad/>')
            with self.subTest(name=name),self.assertRaises(InputError):reader.read_dart_taxonomy(path)
    def test_huge_sparse_row_and_column_rejected(self):
        for payload in (b'<worksheet><row r="100001"/></worksheet>', b'<worksheet><row r="1"><c r="ZZ1"/></row></worksheet>'):
            path = self.book();self.rewrite(path,'xl/worksheets/sheet1.xml',payload)
            with self.assertRaises(InputError):reader.read_dart_taxonomy(path)
    def test_wrong_extension_and_invalid_zip(self):
        path = self.root/'bad.xlsx';path.write_bytes(b'not zip')
        with self.assertRaises(InputError):reader.read_dart_taxonomy(path)
        with self.assertRaises(InputError):reader.read_dart_taxonomy(self.root/'taxonomy.xml')
    def test_numeric_types_are_explicit(self):
        self.assertEqual(reader._kind('custom:inventedItemType','xbrli:item'), ('other',False))
        self.assertEqual(reader._kind('xbrli:sharesItemType','xbrli:item'), ('shares',True))
        self.assertEqual(reader._kind('xbrli:stringItemType','xbrldt:dimensionItem'), ('dimension',False))


if __name__ == '__main__':unittest.main()

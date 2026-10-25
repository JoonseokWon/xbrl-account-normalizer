from copy import deepcopy
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
from application.configuration import default_profile,default_catalog,read_json
from application.loading import load_session
from application.reviews import review_rows,apply_review
from application.audit_export import prepare_audit_tables,export_audit_input
from core.audit_contract import SCHEMA
from core.audit_validation import validate_audit_tables
from core.errors import InputError

ROOT=Path(__file__).resolve().parents[1]
INPUTS=ROOT/'samples/새회사변환예제'

def session(confirmed=True):
    profile=default_profile();profile['sheets']['gl']=['GL_Data'];profile['sheets']['fs']=[]
    profile['columnMappings']=read_json(INPUTS/'예제_열연결.json');profile['sourceSystem']='ERP-DEMO-B'
    result=load_session(INPUTS/'01_노바부품_원본원장.xlsx',profile,default_catalog(),False)
    if confirmed:
        settings=read_json(INPUTS/'예제_계정연결.json')
        for row in review_rows(result):apply_review(result,row['key'],settings['mapping'][row['originalCode']],'확인',settings['reviewer'])
    return result

class AuditExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.session=session();cls.support=INPUTS/'02_노바부품_시산표_통제자료.xlsx'
        cls.tables,cls.source,cls.summary=prepare_audit_tables(cls.session,cls.support)
    def test_totals_bank_merge_and_reverse_revenue(self):
        self.assertEqual(self.summary['ledgerRows'],28)
        self.assertEqual(self.summary['debitTotal'],26700000)
        self.assertEqual(self.summary['creditTotal'],26700000)
        bank=next(row for row in self.session.result['movements'] if row['standardCode']=='A1102')
        self.assertEqual(bank['netMovement'],9900000)
        self.assertEqual(len(bank['originalCodes']),2)
        revenue=next(row for row in self.session.result['movements'] if row['standardCode']=='R4101')
        self.assertEqual((revenue['debit'],revenue['credit']),(500000,13400000))
    def test_pending_mapping_blocks_target_file(self):
        with self.assertRaises(InputError):prepare_audit_tables(session(False),self.support)
    def test_missing_control_is_not_synthesized(self):
        tables=deepcopy(self.tables);tables['독립통제'].pop()
        with self.assertRaises(InputError):validate_audit_tables(tables)
    def test_mixed_company_blocks_export(self):
        tables=deepcopy(self.tables);tables['독립통제'][0]['회사코드']='OTHER'
        with self.assertRaises(InputError):validate_audit_tables(tables)
    def test_control_mismatch_is_not_overwritten(self):
        tables=deepcopy(self.tables);tables['독립통제'][0]['차변합계']+=1
        with self.assertRaises(InputError):validate_audit_tables(tables)
    def test_naive_control_timestamp_is_rejected(self):
        tables=deepcopy(self.tables);tables['독립통제'][0]['추출일시']='2026-02-02T09:00:00'
        with self.assertRaises(InputError):validate_audit_tables(tables)
    def test_unbalanced_tb_is_rejected(self):
        tables=deepcopy(self.tables);tables['시산표'][0]['기말잔액']+=1
        with self.assertRaises(InputError):validate_audit_tables(tables)
    def test_conflicting_supplement_mapping_is_rejected(self):
        support=deepcopy(self.tables);support.pop('총계정원장');support['계정매핑'][0]['리드코드']='WRONG'
        with patch('application.audit_export.read_tables',return_value=(support,self.source)):
            with self.assertRaises(InputError):prepare_audit_tables(self.session,self.support)
    def test_saved_schema_identifiers_and_new_output_folder(self):
        before=self.support.read_bytes()
        with tempfile.TemporaryDirectory() as folder:
            first=export_audit_input(self.session,self.support,folder)
            second=export_audit_input(self.session,self.support,folder)
            self.assertNotEqual(first,second)
            wb=load_workbook(first/'통합감사_원장입력.xlsx',read_only=True,data_only=False)
            try:
                self.assertEqual(wb.sheetnames,list(SCHEMA))
                for name,headers in SCHEMA.items():self.assertEqual(list(next(wb[name].values)),headers)
                row=list(wb['총계정원장'].values)[1]
                self.assertEqual(row[2],'001');self.assertTrue(row[4].startswith('00'))
                self.assertTrue(all(cell.data_type!='f' for sheet in wb for cells in sheet for cell in cells))
            finally:wb.close()
        self.assertEqual(self.support.read_bytes(),before)

if __name__=='__main__':unittest.main()

from pathlib import Path
from datetime import date, datetime
import hashlib
from core.errors import InputError
from core.constants import MAPPING_HEADERS
from field_mapping import resolve_fields
from .package_guard import check_package

def read_input(path, profile, inspect_only=False):
    from openpyxl import load_workbook
    path=Path(path)
    check_package(path)
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    source = {'file':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()}
    try:
        if inspect_only:
            return {'file':path.name,'sheets':[
                {'name':ws.title,'firstRowHeaders':list(next(ws.iter_rows(min_row=1,max_row=1,values_only=True),()))}
                for ws in wb]}
        selected = {}
        for role in ('gl', 'fs'):
            names = profile['sheets'][role]
            names = [names] if isinstance(names, str) else names
            matches = [name for name in names if name in wb.sheetnames]
            if len(matches) > 1:
                raise InputError(f'{role} 입력 시트가 여러 개입니다. 프로필에서 하나를 지정하세요.')
            if matches:
                selected[role] = matches[0]
        if not selected:
            raise InputError('원장 또는 재무제표 시트가 없습니다. config/profile.json의 시트명·열 이름을 원본에 맞추세요.')
        if len(set(selected.values())) != len(selected):
            raise InputError('원장·재무제표 시트는 서로 달라야 합니다.')
        mapping_name = profile['sheets']['mapping']
        if mapping_name in selected.values():
            raise InputError('매핑표와 원본 시트는 달라야 합니다.')
        if mapping_name in wb.sheetnames:
            selected['mapping'] = mapping_name
        raw, matrices, mapping = {}, {}, []
        for role, name in selected.items():
            ws = wb[name]
            ws.reset_dimensions()
            header_row = int(profile.get('headerRows', {}).get(role, 1))
            if not 1 <= header_row <= 100:
                raise InputError('헤더 행은 1~100행 범위가 필요합니다.')
            matrix, records, headers = [], [], None
            for line_no, cells in enumerate(ws.iter_rows(), 1):
                if len(cells) > 80 or line_no > 50001:
                    raise InputError('시트 행·열 한도 초과')
                values = []
                for cell in cells:
                    if cell.data_type in ('f', 'e'):
                        raise InputError(f'{name}!{cell.coordinate}: 수식·오류 셀을 확정값으로 준비하세요.')
                    value = cell.value
                    if isinstance(value, datetime) and value.time() == datetime.min.time():
                        value = value.date().isoformat()
                    elif isinstance(value, date) and not isinstance(value, datetime):
                        value = value.isoformat()
                    if isinstance(value, str) and len(value) > 2000:
                        raise InputError(f'{name}!{cell.coordinate}: 텍스트 길이 초과')
                    values.append(value)
                while values and values[-1] is None:
                    values.pop()
                if line_no < header_row:
                    continue
                if line_no == header_row:
                    if not values or any(not isinstance(h, str) or not h.strip() for h in values) or len(set(values)) != len(values):
                        raise InputError(f'{name}: 빈 헤더·중복 헤더')
                    headers = values
                    matrix.append(values)
                    continue
                if not values or all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                    continue
                if len(values) > len(headers):
                    raise InputError(f'{name}!{line_no}: 헤더 밖 값')
                values += [None] * (len(headers) - len(values))
                matrix.append(values)
                records.append(dict(zip(headers, values)) | {'_source':f'{path.name} / {name}!{line_no}'})
            if headers is None:
                raise InputError(f'{name}: 헤더 행이 없습니다.')
            matrices[role] = matrix
            if role == 'mapping':
                if headers != MAPPING_HEADERS:
                    raise InputError('매핑표의 열 이름·순서가 지원 규격과 다릅니다.')
                mapping = [{k:row[k] for k in MAPPING_HEADERS} for row in records]
                continue
            fields = resolve_fields(headers, role, profile)
            required = {'account','name','date'} | ({'journal','line','debit','credit'} if role == 'gl' else {'amount','unit','rowType','statement'})
            if not profile.get('companyCode'):
                required.add('company')
            if required - fields.keys():
                raise InputError(f'{name}: 필요한 열 누락: {", ".join(sorted(required-fields.keys()))}')
            raw[role] = [{key:row[header] for key,header in fields.items()} | {'_source':row['_source']} for row in records]
        if not any(raw.values()):
            raise InputError('변환할 원본 데이터 행이 없습니다.')
        source['ignoredSheets'] = [name for name in wb.sheetnames if name not in selected.values()]
        return raw, mapping, matrices, source
    finally:
        wb.close()

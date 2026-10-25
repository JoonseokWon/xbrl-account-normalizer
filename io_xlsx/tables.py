"""Strict value-only supplementary tables; no formula cache or external links."""
from datetime import date,datetime
import hashlib
from pathlib import Path
from openpyxl import load_workbook
from core.errors import InputError
from .package_guard import check_package


def read_tables(path,schema,optional=None):
    path=Path(path);check_package(path)
    raw=path.read_bytes();source={'file':path.name,'sha256':hashlib.sha256(raw).hexdigest(),'bytes':len(raw)}
    wb=load_workbook(path,read_only=True,data_only=False,keep_links=False)
    try:
        missing=set(schema)-set(wb.sheetnames)
        if missing:raise InputError('보조파일 필수 시트 누락: '+', '.join(sorted(missing)))
        expected=dict(schema)|{k:v for k,v in (optional or {}).items() if k in wb.sheetnames}
        if set(wb.sheetnames)-set(expected):raise InputError('보조파일에는 지원하는 시트만 포함해 주세요.')
        tables={}
        for name,headers in expected.items():
            sheet=wb[name];sheet.reset_dimensions();rows=[]
            for number,cells in enumerate(sheet.iter_rows(),1):
                if number>10001 or len(cells)>80:raise InputError('통합감사 입력은 시트당 10,000행·80열 이내여야 합니다.')
                values=[]
                for cell in cells:
                    if cell.data_type in ('f','e'):raise InputError(f'{name}!{cell.coordinate}: 수식·오류 대신 확정값이 필요합니다.')
                    value=cell.value
                    if isinstance(value,datetime):value=value.date().isoformat() if value.time()==datetime.min.time() else value.isoformat()
                    elif isinstance(value,date):value=value.isoformat()
                    if isinstance(value,str) and len(value)>2000:raise InputError('텍스트 길이 초과')
                    values.append(value)
                while values and values[-1] is None:values.pop()
                if number==1:
                    if values!=headers:raise InputError(name+': 열 이름·순서를 확인해 주세요.')
                    continue
                if not values or all(v in (None,'') for v in values):continue
                if len(values)>len(headers):raise InputError(name+': 헤더 밖 값')
                values += [None]*(len(headers)-len(values));rows.append(dict(zip(headers,values)))
            tables[name]=rows
        if hashlib.sha256(path.read_bytes()).hexdigest()!=source['sha256']:raise InputError('읽는 중 보조파일이 변경되었습니다.')
        return tables,source
    finally:wb.close()

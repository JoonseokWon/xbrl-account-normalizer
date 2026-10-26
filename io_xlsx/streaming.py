"""Write large review workbooks without retaining an Excel cell graph in memory."""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill
from core.constants import MAPPING_HEADERS

def export_streaming(path,result,matrices,catalog):
    wb=Workbook(write_only=True);header_font=Font(name='맑은 고딕',size=11,bold=True,color='FFFFFF');body_font=Font(name='맑은 고딕',size=10);fill=PatternFill('solid',fgColor='203B5A')
    def add(name,headers,rows):
        sheet=wb.create_sheet(name);sheet.freeze_panes='A2';sheet.sheet_view.showGridLines=False
        for col in range(len(headers)):
            from openpyxl.utils import get_column_letter
            sheet.column_dimensions[get_column_letter(col+1)].width=24
        title=[]
        for value in headers:
            cell=WriteOnlyCell(sheet,value);cell.data_type='s';cell.font=header_font;cell.fill=fill;title.append(cell)
        sheet.append(title)
        for row in rows:
            output=[]
            for value in row:
                cell=WriteOnlyCell(sheet,value)
                cell.font=body_font
                if isinstance(value,str):cell.data_type='s'
                elif isinstance(value,(int,float)):cell.number_format='#,##0;(#,##0);0'
                output.append(cell)
            sheet.append(output)
    add('변환결과',['항목','값'],[['버전',result['version']],['상태',result['status']],['원본 원장 행',result['counts']['sourceGlRows']],['변환 원장 행',len(result['ledger'])],['미확정 행',result['counts']['pendingRows']],['오류',result['counts']['errors']]])
    fields=['company','journalId','lineId','postingDate','originalCode','originalName','standardCode','standardName','debit','credit','mappingStatus','source']
    add('표준원장',['회사코드','전표ID','행ID','전기일','원본계정코드','원본계정명','표준계정코드','표준계정명','차변','대변','상태','원본출처'],([r.get(k) for k in fields] for r in result['ledger']))
    add('전표별분개',['회사','전표','행','차변 계정','차변','대변 계정','대변','원본출처'],([r['company'],r['journalId'],r['lineId'],r['originalName'] if r['debit'] else '',r['debit'],r['originalName'] if r['credit'] else '',r['credit'],r['source']] for r in result['ledger']))
    add('세부정보',['회사','전표','행','구분','값','원본출처'],([r['company'],r['journalId'],r['lineId'],k,v,r['source']] for r in result['ledger'] for k,v in r.get('dimensions',{}).items()))
    add('변환대사',['항목','원본','변환','차이','상태','원본출처'],([r[k] for k in ('label','source','converted','difference','status','reference')] for r in result.get('reconciliation',[])))
    add('표준계정집계',['회사','표준코드','계정명','차변','대변','순변동'],([r[k] for k in ('company','standardCode','standardName','debit','credit','netMovement')] for r in result['movements']))
    add('표준재무제표',['회사','원본코드','원본계정명','표준코드','금액','기준일','원본출처'],([r[k] for k in ('company','originalCode','originalName','standardCode','amount','asOfDate','source')] for r in result['financial']))
    add('매핑표',MAPPING_HEADERS,([r.get(k) for k in MAPPING_HEADERS] for r in result['mappings']))
    add('변환이슈',['구분','코드','내용','원본출처'],([r[k] for k in ('severity','code','message','source')] for r in result['issues']))
    add('표준계정사전',['코드','계정명','종류','상위코드','리드','구분','분석대상'],([r.get(k) for k in ('code','name','kind','parent','lead','statement','module')] for r in catalog['accounts']))
    for role,name in (('gl','원장'),('fs','재무제표')):
        if matrices.get(role):add(name,matrices[role][0],iter(matrices[role][1:]))
    with Path(path).open('xb') as output:wb.save(output)

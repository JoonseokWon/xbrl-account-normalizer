"""Value-only streaming audit input. Rows are neither split nor truncated."""
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font,PatternFill
from openpyxl.utils import get_column_letter

def write_audit(path,tables,schema,widths):
    workbook=Workbook(write_only=True);body=Font(name='맑은 고딕',size=10)
    header=Font(name='맑은 고딕',size=10,bold=True,color='FFFFFF');fill=PatternFill('solid',fgColor='203B5A')
    for name,fields in schema.items():
        sheet=workbook.create_sheet(name);sheet.sheet_view.showGridLines=False;sheet.freeze_panes='A2'
        for i,field in enumerate(fields,1):sheet.column_dimensions[get_column_letter(i)].width=widths.get(field,20)
        title=[]
        for field in fields:
            cell=WriteOnlyCell(sheet,field);cell.data_type='s';cell.font=header;cell.fill=fill;title.append(cell)
        sheet.append(title)
        for row in tables[name]:
            cells=[]
            for field in fields:
                value=row[field];cell=WriteOnlyCell(sheet,value);cell.font=body
                if isinstance(value,str):cell.data_type='s'
                elif isinstance(value,(int,float)):cell.number_format='#,##0;(#,##0);0'
                cells.append(cell)
            sheet.append(cells)
    with path.open('xb') as stream:workbook.save(stream)

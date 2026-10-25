from pathlib import Path
from openpyxl import load_workbook
from core.errors import InputError
from .package_guard import check_package

def inspect_workbook(path,header_row=1):
    path=Path(path)
    check_package(path)
    if not 1<=int(header_row)<=100:
        raise InputError('제목 행은 1~100행으로 입력해 주세요.')
    wb=load_workbook(path,read_only=True,data_only=False,keep_links=False)
    try:
        sheets=[]
        for ws in wb:
            ws.reset_dimensions()
            cells=next(ws.iter_rows(min_row=header_row,max_row=header_row),())
            headers=[str(cell.value).strip() if cell.value is not None else '' for cell in cells]
            while headers and not headers[-1]: headers.pop()
            sheets.append({'name':ws.title,'headers':headers})
        return sheets
    finally:
        wb.close()

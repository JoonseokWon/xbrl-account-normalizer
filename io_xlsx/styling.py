from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo

def make_sheet_builder(wb):
    table_counter=0
    def add(name, headers, rows, editable=False, widths=None):
            nonlocal table_counter
            ws=wb.create_sheet(name)
            ws.sheet_view.showGridLines=False
            ws.append(headers)
            for row in rows:
                ws.append(row)
            for row in ws.iter_rows():
                for cell in row:
                    # Force original strings to text; spreadsheet-like descriptions
                    # must not turn into formulas in the generated workbook.
                    if isinstance(cell.value,str):
                        cell.data_type='s'
                    cell.font=Font(name='맑은 고딕',size=10,color='172B4D')
                    cell.alignment=Alignment(vertical='center')
                    if cell.row>1 and isinstance(cell.value,(int,float)):
                        cell.number_format='#,##0;(#,##0);0'
                    if editable and cell.row>1:
                        cell.fill=PatternFill('solid',fgColor='FFF4CC')
                ws.row_dimensions[row[0].row].height=24
            for cell in ws[1]:
                cell.fill=PatternFill('solid',fgColor='203B5A')
                cell.font=Font(name='맑은 고딕',size=10,bold=True,color='FFFFFF')
                cell.alignment=Alignment(horizontal='center',vertical='center')
            from openpyxl.utils import get_column_letter
            for column,header in enumerate(headers,1):
                ws.column_dimensions[get_column_letter(column)].width=(widths or {}).get(header, max(17,min(35,len(header)*2+4)))
            ws.freeze_panes='A2'
            if rows:
                table_counter+=1
                table=Table(displayName=f'AccountData{table_counter}',ref=f'A1:{get_column_letter(len(headers))}{len(rows)+1}')
                table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=False)
                ws.add_table(table)
            return ws
    return add

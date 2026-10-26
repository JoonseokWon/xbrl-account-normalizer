from pathlib import Path
from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from core.constants import MAPPING_HEADERS
from .styling import make_sheet_builder

def export_review(path, result, matrices, catalog):
    if result['counts']['sourceGlRows']+result['counts']['sourceFsRows']>10000:
        from .streaming import export_streaming
        return export_streaming(path,result,matrices,catalog)
    wb=Workbook()
    wb.remove(wb.active)
    add=make_sheet_builder(wb)
    add('변환결과',['항목','값','설명'],[
        ['버전',result['version'],result['validationStatus']],
        ['처리 상태',result['status'],'계정 매핑 상태이며 감사 수행 완료가 아님'],
        ['원장 원본 행',result['counts']['sourceGlRows'],'원본 행 위치·계정코드 보존'],
        ['원장 변환 행',result['counts']['normalizedGlRows'],'오류 행은 원본 시트와 변환이슈에서 확인'],
        ['재무상태표 변환 행',result['counts']['normalizedFsRows'],'기말잔액이며 원장 당기 변동액과 구분'],
        ['미확정 행',result['counts']['pendingRows'],'매핑표의 표준계정코드·상태·확인자를 입력하고 재변환'],
        ['오류 건수',result['counts']['errors'],'변환이슈 확인'],
        ['원본 파일',result['source']['file'],'입력 파일은 덮어쓰지 않음'],
        ['후보 점수','문자열 유사도','정확도·확률이 아니며 점수만으로 자동 승인하지 않음'],
        ['이력 재사용','같은 회사·시스템·계정코드·계정명·사전 버전','명시적으로 미확정으로 바꾼 매핑은 자동 적용하지 않음'],
        ['집계 범위','확인된 원장행만','미확정 행은 표준원장에 남고 집계에서는 제외'],
    ],widths={'항목':27,'값':55,'설명':80})
    fields=[('company','회사코드'),('originalCode','원본계정코드'),('originalName','원본계정명'),
            ('standardCode','표준계정코드'),('standardName','표준계정명'),('leadCode','리드코드'),
            ('mappingStatus','상태'),('mappingBasis','매핑근거'),('postingDate','전기일'),
            ('journalId','전표ID'),('lineId','행ID'),('debit','차변_원'),('credit','대변_원'),
            ('netMovement','당기순변동_원'),('module','분석대상'),('source','원본출처')]
    add('표준원장',[h for k,h in fields],[[r.get(k) for k,h in fields] for r in result['ledger']],widths={'매핑근거':55,'원본출처':70,'표준계정명':30})
    add('전표별분개',['회사','전표','행','전기일','차변 계정','차변','대변 계정','대변','원본출처'],
        [[r['company'],r['journalId'],r['lineId'],r['postingDate'],r['originalName'] if r['debit'] else '',r['debit'],r['originalName'] if r['credit'] else '',r['credit'],r['source']] for r in result['ledger']])
    add('세부정보',['회사','전표','행','구분','값','원본출처'],[[r['company'],r['journalId'],r['lineId'],k,v,r['source']] for r in result['ledger'] for k,v in r.get('dimensions',{}).items()])
    add('변환대사',['확인 항목','원본 / 시산표','변환 원장','차이','상태','원본출처'],[[r[k] for k in ('label','source','converted','difference','status','reference')] for r in result.get('reconciliation',[])])
    group_fields=['company','standardCode','standardName','leadCode','originalCodes','rowCount','debit','credit','netMovement']
    add('표준계정집계',['회사코드','표준계정코드','표준계정명','리드코드','원본계정코드들','원장행수','차변_원','대변_원','당기순변동_원'],
        [[', '.join(r[k]) if k=='originalCodes' else r[k] for k in group_fields] for r in result['movements']])
    fs_fields=['company','originalCode','originalName','standardCode','standardName','rowType','amount','asOfDate','mappingStatus','mappingBasis','source']
    add('표준재무제표',['회사코드','원본항목코드','원본항목명','표준계정코드','표준계정명','행구분','기말잔액_원','기준일','상태','매핑근거','원본출처'],
        [[r.get(k) for k in fs_fields] for r in result['financial']],widths={'매핑근거':55,'원본출처':70})
    unique_candidates={}
    for kind,rows in [('원장',result['ledger']),('재무제표',result['financial'])]:
        for row in rows:
            key=(kind,row['company'],row['system'],row['originalCode'],row['originalName'])
            if key in unique_candidates:
                continue
            unique_candidates[key]=True
            for rank,candidate in enumerate(row['candidateDetails'],1):
                unique_candidates[(key,rank)]=[*key,rank,candidate['code'],candidate['name'],candidate['score'],candidate['reason'],candidate['matchedLabel']]
    candidate_rows=[v for v in unique_candidates.values() if isinstance(v,list)]
    candidates=add('매핑후보',['입력종류','회사코드','원천시스템','원본계정코드','원본계정명','순위','후보코드','후보계정명','유사도','추천근거','일치한별칭'],candidate_rows,
                   widths={'추천근거':47,'일치한별칭':32,'후보계정명':32})
    for row in range(2,candidates.max_row+1):
        candidates.cell(row,9).number_format='0.0000'
    mapping=add('매핑표',MAPPING_HEADERS,[[r.get(h) for h in MAPPING_HEADERS] for r in result['mappings']],True)
    choices=DataValidation(type='list',formula1=f"'표준계정사전'!$A$2:$A${len(catalog['accounts'])+1}")
    choices.errorTitle='표준계정코드 확인'; choices.error='표준계정사전의 코드를 선택하세요.'; choices.showErrorMessage=True
    mapping.add_data_validation(choices);choices.add(f'F2:F{max(mapping.max_row,2)}')
    states=DataValidation(type='list',formula1='"미확정,확인,거절"')
    mapping.add_data_validation(states);states.add(f'G2:G{max(mapping.max_row,2)}')
    add('변환이슈',['구분','이슈코드','내용','원본출처'],[[r['severity'],r['code'],r['message'],r['source']] for r in result['issues']],widths={'내용':85,'원본출처':70,'이슈코드':31})
    add('표준계정사전',['표준계정코드','계정명','계정종류','상위코드','리드코드','표시구분','분석대상'],
        [[r['code'],r['name'],r['kind'],r.get('parent'),r['lead'],r['statement'],r['module']] for r in catalog['accounts']],widths={'계정명':35,'분석대상':30})
    for role,name in [('gl','원장'),('fs','재무제표')]:
        if matrices.get(role):
            add(name,matrices[role][0],matrices[role][1:])
    path=Path(path)
    if path.exists():
        raise FileExistsError('기존 검토 엑셀을 덮어쓸 수 없습니다.')
    # Exclusive creation prevents a second process from replacing an output.
    with path.open('xb') as stream:
        wb.save(stream)

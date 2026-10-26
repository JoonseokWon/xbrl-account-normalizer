"""Value-only arbitrary source table, shared by supplementary input mapping."""
from datetime import date,datetime
import hashlib
from pathlib import Path
from core.errors import InputError
from core.limits import MAX_ROWS,MAX_COLUMNS
from .package_guard import check_package
from .delimited import read_delimited

def read_records(path,sheet,header_row=1):
    path=Path(path)
    if path.suffix.lower() in ('.csv','.tsv'):return read_delimited(path,header_row)
    from openpyxl import load_workbook
    check_package(path)
    if not 1<=header_row<=100:raise InputError('제목 행은 1~100행이어야 합니다.')
    digest=hashlib.sha256(path.read_bytes()).hexdigest()
    wb=load_workbook(path,read_only=True,data_only=False,keep_links=False)
    try:
        if sheet not in wb.sheetnames:raise InputError('선택한 시트를 찾을 수 없습니다: '+str(sheet))
        ws=wb[sheet];ws.reset_dimensions();records=[];matrix=[];headers=None
        for number,cells in enumerate(ws.iter_rows(),1):
            if number<header_row:continue
            if len(cells)>MAX_COLUMNS:raise InputError('80열 한도를 초과했습니다.')
            values=[]
            for cell in cells:
                if cell.data_type in ('f','e'):raise InputError(f'{sheet}!{cell.coordinate}: 수식·오류 대신 확정값이 필요합니다.')
                value=cell.value
                if isinstance(value,datetime):value=value.date().isoformat() if value.time()==datetime.min.time() else value.isoformat()
                elif isinstance(value,date):value=value.isoformat()
                if isinstance(value,str) and len(value)>2000:raise InputError('셀당 2,000자 한도를 초과했습니다.')
                values.append(value)
            while values and values[-1] is None:values.pop()
            if number==header_row:
                if not values or any(not isinstance(v,str) or not v.strip() for v in values) or len(set(values))!=len(values):raise InputError('빈 제목·중복 제목을 확인해 주세요.')
                headers=values;matrix.append(values);continue
            if not values or all(v in (None,'') for v in values):continue
            if len(values)>len(headers):raise InputError('제목 밖 값이 있습니다.')
            values += [None]*(len(headers)-len(values));matrix.append(values)
            records.append(dict(zip(headers,values))|{'_source':f'{path.name} / {sheet}!{number}'})
            if len(records)>MAX_ROWS:raise InputError('파일당 데이터는 100,000행 이내여야 합니다.')
        if headers is None:raise InputError('제목 행이 없습니다.')
        if hashlib.sha256(path.read_bytes()).hexdigest()!=digest:raise InputError('읽는 중 원본 파일이 변경되었습니다.')
        return records,matrix,{'file':path.name,'sha256':digest,'sheet':sheet,'headerRow':header_row}
    finally:wb.close()

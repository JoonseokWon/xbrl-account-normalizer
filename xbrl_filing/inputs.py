"""Bounded, value-only financial-report inputs and unsigned review work files."""
from copy import deepcopy
import csv
from datetime import date,datetime,time
from decimal import Decimal,InvalidOperation,localcontext
import hashlib
from io import BytesIO,StringIO
import json
import os
from pathlib import Path,PurePosixPath
import re
import tempfile
from xml.parsers import expat
from zipfile import ZipFile,BadZipFile

from core.errors import InputError

SCHEMA='account-normalizer-filing/1'
MAX_BYTES=50*1024**2
MAX_EXPANDED=256*1024**2
MAX_FACTS=50_000
MAX_TEXT=100_000
HEADERS=['항목ID','재무제표','원본코드','원본항목','공식개념','값','단위','표시배율','정확도','기간형식','기준일','시작일','종료일','차원JSON']
INFO=['회사명','기업식별자','식별자체계URI','당기시작일','당기종료일','연결구분','비교기간JSON']
UNITS={'KRW','shares','KRW/shares','pure',''}


def _text(value,label,blank=True,limit=2000):
    if value is None and blank:return ''
    if not isinstance(value,str) or len(value)>limit or (not blank and not value.strip()):
        raise InputError(label+': 텍스트 형식·길이를 확인해 주세요. 코드의 앞자리 0은 텍스트로 보관해 주세요.')
    if any(not (ord(c) in (9,10,13) or 32<=ord(c)<=0xD7FF or 0xE000<=ord(c)<=0xFFFD or 0x10000<=ord(c)<=0x10FFFF) for c in value):
        raise InputError(label+': 지원하지 않는 제어 문자가 있습니다.')
    return value.strip()


def _day(value,label,blank=True):
    if value is None or value=='':
        if blank:return ''
        raise InputError(label+': 날짜를 입력해 주세요.')
    if isinstance(value,datetime):
        if value.time()!=time():raise InputError(label+': 시각 없이 날짜만 입력해 주세요.')
        value=value.date()
    if isinstance(value,date):return value.isoformat()
    text=_text(value,label,False)
    try:
        if re.fullmatch(r'\d{4}-\d{2}-\d{2}',text,flags=re.ASCII) is None:raise ValueError()
        return date.fromisoformat(text).isoformat()
    except ValueError as exc:raise InputError(label+': YYYY-MM-DD 또는 Excel 날짜 셀이 필요합니다.') from exc


def _duplicate_free(pairs):
    result={}
    for key,value in pairs:
        if key in result:raise InputError('JSON에 중복된 키가 있습니다: '+str(key)[:100])
        result[key]=value
    return result


def _json(text):
    def nonfinite(value):raise InputError('JSON에는 NaN·Infinity를 사용할 수 없습니다.')
    try:return json.loads(text,object_pairs_hook=_duplicate_free,parse_constant=nonfinite)
    except InputError:raise
    except (ValueError,RecursionError) as exc:raise InputError('정상적인 UTF-8 JSON 구조를 확인해 주세요.') from exc


def _bounded(value):
    stack=[(value,0)];count=0;text_size=0
    while stack:
        item,depth=stack.pop();count+=1
        if depth>24 or count>2_000_000:raise InputError('공시 JSON의 중첩·항목 수가 한도를 넘었습니다.')
        if isinstance(item,str):
            _text(item,'공시 텍스트',limit=MAX_TEXT);text_size+=len(item)
            if text_size>20_000_000:raise InputError('공시 입력 텍스트 합계가 한도를 넘었습니다.')
        elif isinstance(item,dict):
            for key,child in item.items():
                _text(key,'JSON 키',False);stack.append((child,depth+1))
        elif isinstance(item,list):stack.extend((child,depth+1) for child in item)
        elif item is not None and not isinstance(item,(bool,int,float)):raise InputError('공시 JSON 값 형식을 확인해 주세요.')
        elif isinstance(item,float) and not Decimal(str(item)).is_finite():raise InputError('유한한 숫자만 입력해 주세요.')


def _snapshot(path):
    path=Path(path).resolve()
    if not path.is_file() or path.stat().st_size>MAX_BYTES:raise InputError('50MB 이하의 공시 입력 파일을 선택해 주세요.')
    with path.open('rb') as stream:data=stream.read(MAX_BYTES+1)
    if len(data)>MAX_BYTES:raise InputError('공시 입력 파일은 50MB 이하여야 합니다.')
    return path,data,hashlib.sha256(data).hexdigest()


def _unchanged(path,digest):
    current=hashlib.sha256();size=0
    with path.open('rb') as stream:
        while block:=stream.read(1024**2):
            size+=len(block)
            if size>MAX_BYTES:raise InputError('읽는 중 원본이 변경되었습니다.')
            current.update(block)
    if current.hexdigest()!=digest:raise InputError('읽는 중 원본이 변경되었습니다. 다시 열어 주세요.')


def _package(data):
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries=archive.infolist()
            if len(entries)>1000 or sum(e.file_size for e in entries)>MAX_EXPANDED:raise InputError('공시 엑셀 압축해제 크기·파일 수 한도를 넘었습니다.')
            seen=set();total_text=[0]
            for entry in entries:
                name=entry.filename;pure=PurePosixPath(name);key=str(pure).casefold()
                if key in seen or '\\' in entry.orig_filename or any(ord(c)<32 for c in entry.orig_filename) or pure.is_absolute() or '..' in pure.parts or ':' in name or entry.flag_bits&1 or entry.file_size>128*1024**2 or (entry.external_attr>>16)&0o170000==0o120000:
                    raise InputError('공시 엑셀에 중복·외부 경로 등 안전하지 않은 압축 항목이 있습니다.')
                seen.add(key);lower=name.lower()
                if any(s in lower for s in ('externallinks/','vba','embeddings/','connections.xml')):raise InputError('매크로·외부 연결·내장 개체가 있는 엑셀은 지원하지 않습니다.')
                if not lower.endswith(('.xml','.rels')):continue
                parser=expat.ParserCreate(namespace_separator='}');depth=[0];text_node=[0]
                def forbidden(*unused):raise InputError('DOCTYPE·XML 엔터티가 있는 공시 엑셀은 지원하지 않습니다.')
                def start(tag,attrs):
                    depth[0]+=1
                    text_node[0]=0
                    if depth[0]>40:raise InputError('공시 엑셀 XML 중첩 한도를 넘었습니다.')
                    local=tag.rsplit('}',1)[-1]
                    if local=='Relationship' and attrs.get('TargetMode','').casefold()=='external':raise InputError('외부 링크가 있는 공시 엑셀은 지원하지 않습니다.')
                    if any('macroenabled' in str(v).casefold() for v in attrs.values()):raise InputError('매크로 통합문서는 지원하지 않습니다.')
                    if lower.startswith('xl/worksheets/'):
                        if local=='f':raise InputError('수식 대신 확정값을 입력해 주세요.')
                        if local=='row':
                            index=attrs.get('r','')
                            if not index.isascii() or not index.isdigit() or not 1<=int(index)<=MAX_FACTS+1:raise InputError('공시 시트는 50,000개 데이터 행 이내여야 합니다.')
                        if local=='c':
                            coordinate=attrs.get('r','');match=re.fullmatch(r'([A-Z]+)[1-9][0-9]*',coordinate)
                            if not match:raise InputError('공시 엑셀 셀 좌표가 올바르지 않습니다.')
                            column=0
                            for char in match[1]:column=column*26+ord(char)-64
                            if column>80 or attrs.get('t')=='e':raise InputError('공시 엑셀 열 한도·오류 셀을 확인해 주세요.')
                def end(tag):depth[0]-=1
                def characters(text):
                    text_node[0]+=len(text);total_text[0]+=len(text)
                    if text_node[0]>MAX_TEXT or total_text[0]>20_000_000:raise InputError('공시 엑셀의 셀·텍스트 합계 한도를 넘었습니다.')
                parser.StartDoctypeDeclHandler=forbidden;parser.EntityDeclHandler=forbidden;parser.ExternalEntityRefHandler=forbidden
                parser.StartElementHandler=start;parser.EndElementHandler=end;parser.CharacterDataHandler=characters
                with archive.open(entry) as stream:
                    while block:=stream.read(65536):parser.Parse(block,False)
                parser.Parse(b'',True)
    except (BadZipFile,expat.ExpatError,RuntimeError,KeyError) as exc:raise InputError('정상적인 공시 .xlsx 파일이 아닙니다.') from exc


def _number(value,scale):
    if isinstance(value,bool) or value is None:raise InputError('공시 숫자 값이 필요합니다.')
    if isinstance(value,str):
        text=value.strip()
        if text.startswith('(') and text.endswith(')'):text='-'+text[1:-1].strip()
        if len(text)>128 or re.fullmatch(r'[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?',text,flags=re.ASCII) is None:raise InputError('금액은 숫자·천단위 쉼표·괄호 음수로 입력해 주세요.')
        text=text.replace(',','')
    elif isinstance(value,(int,float,Decimal)):text=str(value)
    else:raise InputError('숫자 값 형식을 확인해 주세요.')
    try:
        with localcontext() as context:
            context.prec=120;number=Decimal(text)*scale
    except InvalidOperation as exc:raise InputError('숫자 값 형식을 확인해 주세요.') from exc
    if not number.is_finite() or len(number.as_tuple().digits)>100 or abs(number)>Decimal('1e100') or abs(number.as_tuple().exponent)>100:raise InputError('유한한 100자리 이내 숫자가 필요합니다.')
    if not number:return '0'
    result=format(number,'f')
    result=result.rstrip('0').rstrip('.') if '.' in result else result
    if len(result)>128 or len(Decimal(result).as_tuple().digits)>100:raise InputError('확정 숫자 값은 100자리 이내여야 합니다.')
    return result


def _scale(value):
    if value in (None,''):return 1
    if isinstance(value,bool):raise InputError('표시배율은 1, 1000, 1000000 중 하나여야 합니다.')
    if str(value).strip() not in ('1','1000','1000000'):raise InputError('표시배율은 1, 1000, 1000000 중 하나여야 합니다.')
    return int(str(value).strip())


def _row(record,number,source):
    original_id=_text(record.get('항목ID'),'항목ID',limit=100);original_code=_text(record.get('원본코드'),'원본코드')
    statement=_text(record.get('재무제표'),'재무제표',False)
    if statement not in ('BS','PL','CF','CE','NOTE'):raise InputError('재무제표는 BS/PL/CF/CE/NOTE 중 하나를 입력해 주세요.')
    kind=_text(record.get('기간형식'),'기간형식',False)
    if kind not in ('instant','duration'):raise InputError('기간형식은 instant 또는 duration을 입력해 주세요.')
    if kind=='instant':period={'kind':kind,'instant':_day(record.get('기준일'),'기준일',False)}
    else:period={'kind':kind,'start':_day(record.get('시작일'),'시작일',False),'end':_day(record.get('종료일'),'종료일',False)}
    if kind=='duration' and period['start']>period['end']:raise InputError('항목 시작일이 종료일보다 늦습니다.')
    unit=_text(record.get('단위'),'단위');scale=_scale(record.get('표시배율'))
    if unit not in UNITS:raise InputError('단위는 KRW/shares/KRW/shares/pure 또는 빈 값으로 입력해 주세요.')
    original=record.get('값')
    if unit:value=_number(original,scale)
    else:
        if scale!=1:raise InputError('문자·날짜·참거짓 항목에는 표시배율 1을 사용해 주세요.')
        if isinstance(original,(date,datetime)):value=_day(original,'날짜 값',False)
        elif isinstance(original,bool):value='true' if original else 'false'
        else:value=_text(original if original is None or isinstance(original,str) else str(original),'값',limit=MAX_TEXT)
    decimals=record.get('정확도');decimals='' if not unit else '0' if decimals in (None,'') else str(decimals).strip()
    if unit and (decimals!='INF' and (re.fullmatch(r'-?(?:0|[1-9]\d?)',decimals) is None or not -12<=int(decimals)<=12)):raise InputError('정확도는 -12~12 또는 INF를 입력해 주세요.')
    dimensions=_json(_text(record.get('차원JSON'),'차원 JSON',limit=MAX_TEXT) or '{}')
    if not isinstance(dimensions,dict) or len(dimensions)>32 or any(not isinstance(k,str) or not isinstance(v,str) for k,v in dimensions.items()):raise InputError('차원 JSON은 축과 멤버 QName의 객체여야 합니다.')
    return {'id':f'f{number:06d}','originalId':original_id,'originalCode':original_code,'originalName':_text(record.get('원본항목'),'원본항목'),
            'concept':_text(record.get('공식개념'),'공식개념'),'statement':statement,'period':period,'value':value,'unit':unit,'decimals':decimals,
            'dimensions':dimensions,'source':source,'approved':False,'reviewer':'','originalValue':str(original) if original is not None else '',
            'displayScale':scale,'normalizedToBaseUnit':True}


def _draft(info=None):
    info=info or {};comparative=info.get('비교기간JSON') or '[]'
    comparative=_json(comparative) if isinstance(comparative,str) else comparative
    if not isinstance(comparative,list) or len(comparative)>3 or any(not isinstance(p,dict) for p in comparative):raise InputError('비교기간 JSON은 기간 객체 0~3개의 배열이어야 합니다.')
    for period in comparative:
        period['periodStart']=_day(period.get('periodStart'),'비교 시작일',False);period['periodEnd']=_day(period.get('periodEnd'),'비교 종료일',False)
    consolidation=_text(info.get('연결구분'),'연결구분') or '별도'
    if consolidation not in ('별도','연결'):raise InputError('연결구분은 별도 또는 연결을 입력해 주세요.')
    return {'schema':SCHEMA,'entity':{'name':_text(info.get('회사명'),'회사명'),'identifier':_text(info.get('기업식별자'),'기업식별자'),
            'scheme':_text(info.get('식별자체계URI'),'식별자체계URI')},'reporting':{'periodStart':_day(info.get('당기시작일'),'당기시작일'),
            'periodEnd':_day(info.get('당기종료일'),'당기종료일'),'consolidation':consolidation,'comparativePeriods':comparative},'facts':[]}


def _excel(data,name):
    _package(data)
    from openpyxl import load_workbook
    wb=load_workbook(BytesIO(data),read_only=True,data_only=False,keep_links=False)
    try:
        if not {'공시정보','공시항목'}<=set(wb.sheetnames):raise InputError('공시정보·공시항목 시트가 필요합니다. 공시 엑셀 양식을 사용해 주세요.')
        info={};sheet=wb['공시정보'];sheet.reset_dimensions()
        for number,cells in enumerate(sheet.iter_rows(),1):
            values=[c.value for c in cells]
            if number==1:
                if values[:2]!=['항목','값']:raise InputError('공시정보 제목은 항목/값이어야 합니다.')
                continue
            if not values or all(v in (None,'') for v in values):continue
            if len(values)>2 and any(v not in (None,'') for v in values[2:]):raise InputError('공시정보 제목 밖에 값이 있습니다.')
            key=_text(values[0],'공시정보 항목',False)
            if key not in INFO or key in info:raise InputError('공시정보에 알 수 없거나 중복된 항목이 있습니다.')
            info[key]=values[1] if len(values)>1 else None
        result=_draft(info);sheet=wb['공시항목'];sheet.reset_dimensions();headers=None
        for number,cells in enumerate(sheet.iter_rows(),1):
            values=[c.value for c in cells]
            while values and values[-1] is None:values.pop()
            if number==1:
                if values!=HEADERS:raise InputError('공시항목 열 이름·순서를 입력 양식과 맞춰 주세요.')
                headers=values;continue
            if not values or all(v in (None,'') for v in values):continue
            if len(values)>len(HEADERS):raise InputError('공시항목 제목 밖에 값이 있습니다.')
            values += [None]*(len(HEADERS)-len(values))
            result['facts'].append(_row(dict(zip(HEADERS,values)),len(result['facts'])+1,f'{name}:공시항목:{number}'))
            if len(result['facts'])>MAX_FACTS:raise InputError('공시 항목은 50,000개 이내여야 합니다.')
        if headers is None:raise InputError('공시항목 제목 행이 없습니다.')
        return result
    finally:wb.close()


def _csv(data,name):
    try:text=data.decode('utf-8-sig')
    except UnicodeDecodeError:
        try:text=data.decode('cp949')
        except UnicodeDecodeError as exc:raise InputError('CSV를 UTF-8 또는 CP949로 저장해 주세요.') from exc
    previous=csv.field_size_limit();csv.field_size_limit(MAX_TEXT)
    try:
        reader=csv.reader(StringIO(text,newline=''),strict=True)
        if next(reader,None)!=HEADERS:raise InputError('CSV 열 이름·순서를 공시 입력 양식과 맞춰 주세요.')
        result=_draft()
        for number,values in enumerate(reader,2):
            if not values or all(not v.strip() for v in values):continue
            if any(value.lstrip().startswith('=') for value in values):raise InputError('CSV에는 수식 대신 확정값을 입력해 주세요.')
            if len(values)!=len(HEADERS):raise InputError(f'CSV {number}행의 열 수가 다릅니다.')
            result['facts'].append(_row(dict(zip(HEADERS,values)),len(result['facts'])+1,f'{name}:CSV:{number}'))
            if len(result['facts'])>MAX_FACTS:raise InputError('공시 항목은 50,000개 이내여야 합니다.')
        return result
    except csv.Error as exc:raise InputError('공시 CSV 구문·셀 길이를 확인해 주세요.') from exc
    finally:csv.field_size_limit(previous)


def _report(value):
    _bounded(value)
    if not isinstance(value,dict) or value.get('schema')!=SCHEMA or not isinstance(value.get('entity'),dict) or not isinstance(value.get('reporting'),dict) or not isinstance(value.get('facts'),list):raise InputError('지원하는 공시 JSON 형식이 아닙니다.')
    if len(value['facts'])>MAX_FACTS:raise InputError('공시 항목은 50,000개 이내여야 합니다.')
    ids=set()
    for fact in value['facts']:
        if not isinstance(fact,dict):raise InputError('공시 항목 객체가 필요합니다.')
        identifier=_text(fact.get('id'),'항목 ID',False,100)
        if identifier in ids:raise InputError('공시 항목 ID가 중복됐습니다.')
        ids.add(identifier)
        for key in ('originalId','originalCode'):
            if key in fact:_text(fact[key],key)
        if not isinstance(fact.get('value'),str):raise InputError('JSON 공시 값은 확정된 원 단위 문자열로 입력해 주세요.')
        if fact.get('approved',False) not in (True,False) or not isinstance(fact.get('approved',False),bool):raise InputError('승인 상태는 true/false여야 합니다.')
    return deepcopy(value)


def read_input(path):
    path,data,digest=_snapshot(path);suffix=path.suffix.lower()
    if suffix=='.xlsx':result=_excel(data,path.name)
    elif suffix=='.csv':result=_csv(data,path.name)
    elif suffix in ('.json','.acfiling'):
        try:result=_report(_json(data.decode('utf-8-sig')))
        except UnicodeDecodeError as exc:raise InputError('공시 JSON은 UTF-8로 저장해 주세요.') from exc
        for fact in result['facts']:fact.update(approved=False,reviewer='')
    else:raise InputError('공시 입력은 .xlsx/.csv/.json 형식을 사용해 주세요.')
    _bounded(result);_unchanged(path,digest)
    result['inputs']=[{'path':str(path),'name':path.name,'sha256':digest}]
    return result


def read_work(path):
    path,data,digest=_snapshot(path)
    if path.suffix.lower() not in ('.acfiling','.json'):raise InputError('공시 업무 .acfiling 파일을 선택해 주세요.')
    try:result=_report(_json(data.decode('utf-8-sig')))
    except UnicodeDecodeError as exc:raise InputError('공시 업무는 UTF-8 JSON이어야 합니다.') from exc
    _unchanged(path,digest);return result


def write_work(report,path):
    path=Path(path).resolve()
    if path.suffix.lower()!='.acfiling':raise InputError('공시 업무 확장자는 .acfiling을 사용해 주세요.')
    value=_report(report);data=json.dumps(value,ensure_ascii=False,allow_nan=False,indent=2).encode('utf-8')
    if len(data)>MAX_BYTES:raise InputError('공시 업무 파일은 50MB 이하여야 합니다.')
    path.parent.mkdir(parents=True,exist_ok=True);handle,name=tempfile.mkstemp(prefix='.filing-work-',dir=path.parent)
    try:
        with os.fdopen(handle,'wb') as stream:stream.write(data);stream.flush();os.fsync(stream.fileno())
        os.replace(name,path)
    finally:
        if Path(name).exists():Path(name).unlink()
    return path


def create_template(path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    path=Path(path).resolve()
    if path.suffix.lower()!='.xlsx':raise InputError('입력 양식 확장자는 .xlsx를 사용해 주세요.')
    wb=Workbook();info=wb.active;info.title='공시정보';info.append(['항목','값'])
    for key in INFO:info.append([key,'[]' if key=='비교기간JSON' else ''])
    facts=wb.create_sheet('공시항목');facts.append(HEADERS)
    examples=wb.create_sheet('작성예시');examples.append(HEADERS)
    examples.append(['0001','BS','0010','현금및현금성자산','', '1,234','KRW',1000,0,'instant','2026-12-31','','','{}'])
    examples.append(['0002','PL','0040','매출액','', '(12,345.50)','KRW',1,2,'duration','','2026-01-01','2026-12-31','{}'])
    examples.append(['0003','NOTE','','회계정책 설명','', '확정한 주석 내용을 입력하세요.','',1,'','duration','','2026-01-01','2026-12-31','{}'])
    guide=wb.create_sheet('사용법');guide.append(['항목','설명'])
    guidance=[('시작하기','공시정보와 공시항목 시트에 회사별 확정 재무제표를 입력하세요. 공시항목은 빈 양식이며 작성예시는 변환하지 않습니다.'),
        ('공식 개념','사용할 공식 택사노미를 선택한 뒤 앱에서 항목별 개념·기간·차원을 확인하고 승인하세요. 예시 개념은 비워 두었습니다.'),
        ('항목ID·원본코드','앞자리 0은 텍스트 셀로 보관하세요. 항목ID는 사용자 관리용이고 XML 내부 ID는 별도로 생성합니다.'),
        ('표시배율','1 / 1000 / 1000000 중 입력하세요. 값 1,234와 표시배율 1000은 원 단위 1234000으로 변환합니다. 원본값과 배율도 보관합니다.'),
        ('정확도','decimals는 금액 정확도이며 표시배율과 별개입니다. -3은 천원 단위 정확도, INF는 정확한 값입니다. -12~12 또는 INF를 사용하세요.'),
        ('기간','BS 등 시점 개념은 instant·기준일, 기간 개념은 duration·시작일·종료일을 입력합니다. 날짜는 YYYY-MM-DD 또는 Excel 날짜 셀입니다.'),
        ('비교기간JSON','예: [{"periodStart":"2025-01-01","periodEnd":"2025-12-31"}]. 당기 외 비교금액을 입력할 때 공시정보에도 해당 기간을 명시하세요.'),
        ('차원JSON','예: {"공식접두사:축이름":"공식접두사:멤버이름"}. 차원이 없으면 {}를 입력하세요. 실제 택사노미의 명시적 차원을 확인하세요.'),
        ('금액·단위','KRW / shares / KRW/shares / pure 또는 문자·날짜·참거짓의 빈 단위를 사용하세요. 쉼표·괄호 음수는 숫자로 변환하며 수식은 허용하지 않습니다.'),
        ('원본·업무 보관','승인 상태는 앱에서 별도 .acfiling 업무로 저장합니다. 업무는 서명된 증빙이 아니며 원본 입력과 공식 택사노미를 함께 보관하세요.'),
        ('공시 준비 범위','원장 거래 합계에서 기말잔액·손익·현금흐름을 추정하지 않습니다. DART 편집기의 추가 작성·제출·접수 확인은 별도입니다.')]
    for row in guidance:guide.append(row)
    for sheet in wb:
        sheet.freeze_panes='A2';sheet.sheet_view.showGridLines=False;sheet.row_dimensions[1].height=30
        for cell in sheet[1]:cell.fill=PatternFill('solid',fgColor='2769FF');cell.font=Font(name='맑은 고딕',bold=True,color='FFFFFF',size=11);cell.alignment=Alignment(vertical='center',wrap_text=True)
        for row in sheet.iter_rows(min_row=2):
            sheet.row_dimensions[row[0].row].height=44 if sheet.title!='사용법' else 58
            for cell in row:cell.font=Font(name='맑은 고딕',size=11);cell.alignment=Alignment(vertical='top',wrap_text=True)
        if sheet.title in ('공시항목','작성예시'):
            widths=[16,12,18,30,44,24,18,15,13,16,18,18,18,44]
            for index,width in enumerate(widths,1):sheet.column_dimensions[get_column_letter(index)].width=width
            sheet.auto_filter.ref='A1:N1' if sheet.title=='공시항목' else f'A1:N{sheet.max_row}'
            for column in ('A','C'):sheet.column_dimensions[column].number_format='@'
            if sheet.title=='공시항목':
                for formula,col in [('"BS,PL,CF,CE,NOTE"','B'),('"KRW,shares,KRW/shares,pure"','G'),('"1,1000,1000000"','H'),('"instant,duration"','J')]:
                    validation=DataValidation(type='list',formula1=formula,allow_blank=True);sheet.add_data_validation(validation);validation.add(f'{col}2:{col}{MAX_FACTS+1}')
        else:sheet.column_dimensions['A'].width=24;sheet.column_dimensions['B'].width=100 if sheet.title=='사용법' else 70
    for cell in info['B']:cell.number_format='@'
    for sheet in (facts,examples):
        for row in sheet.iter_rows(min_row=2):
            for column in (0,2):row[column].number_format='@'
    buffer=BytesIO();wb.save(buffer);wb.close();path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('xb') as stream:stream.write(buffer.getvalue())
    return path

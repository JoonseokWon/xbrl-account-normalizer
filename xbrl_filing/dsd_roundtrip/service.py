"""Template-preserving spreadsheet-to-DSD drafts; never a submission claim."""
from copy import copy
from decimal import Decimal, InvalidOperation, localcontext
import hashlib
import json
import os
from pathlib import Path
import posixpath
import re
import shutil
import tempfile
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

from io_xlsx.package_guard import check_package
from xbrl_filing.disclosure_compare.dsd import read_dsd, _tag, _tables, MAX_XML_BYTES
from xbrl_filing.official_writer.artifacts import _xml, _inventory, _Budget, _read_member

MAPPING_SCHEMA = 'account-normalizer-dsd-workbook-mapping/1'
MAX_CELLS = 100000
MAX_TEXT = 32767
MAX_WORKBOOK_BYTES = 50 * 1024 * 1024
_COORD = re.compile(r'[A-Z]{1,3}[1-9][0-9]{0,6}')


def _digest(path, limit=250 * 1024 * 1024):
    digest, size = hashlib.sha256(), 0
    with Path(path).open('rb') as stream:
        while block := stream.read(1024 * 1024):
            size += len(block)
            if size > limit:
                raise ValueError('입력 파일 크기 한도를 넘었습니다.')
            digest.update(block)
    return digest.hexdigest()


def _value(value):
    if value is None:
        return ''
    if isinstance(value, bool):
        raise ValueError('논리값은 DSD 공시 값으로 자동 변환할 수 없습니다. 텍스트나 숫자로 준비해 주세요.')
    if isinstance(value, (str, int, float, Decimal)):
        text = str(value)
        if len(text) > MAX_TEXT or any(ord(c) < 32 and c not in '\t\n\r' for c in text):
            raise ValueError('셀 텍스트의 길이 또는 제어문자를 확인해 주세요.')
        if isinstance(value, (int, float, Decimal)) and not Decimal(text).is_finite():
            raise ValueError('유한한 숫자만 DSD에 반영할 수 있습니다.')
        return text
    raise ValueError('날짜·오류 등 특수 셀 값은 공시 표시 텍스트로 명시해 주세요.')


def _load_workbook(path):
    path = Path(path)
    check_package(path)
    book = load_workbook(path, read_only=False, data_only=False, keep_links=False)
    try:
        _restore_exact_numbers(path, book)
    except Exception:
        book.close()
        raise
    return book


def _decimal_value(raw):
    if not isinstance(raw, str) or not raw or len(raw) > 200:
        raise ValueError('정산표 숫자 원문의 길이 한도를 확인해 주세요.')
    try:
        number = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError('정산표 숫자 원문을 읽을 수 없습니다.') from exc
    if (not number.is_finite() or abs(number.as_tuple().exponent) > 1000
            or abs(number.adjusted()) > 1000 or len(format(number, 'f')) > 200):
        raise ValueError('정산표 숫자는 유한한 200자리 이내여야 합니다. 큰 수는 명시적인 텍스트로 준비해 주세요.')
    return number


def _restore_exact_numbers(path, book):
    """Preserve OOXML decimal lexicals instead of binary-float rounding.

    Source sheets and relationship targets come from the already guarded ZIP,
    never the network. Provenance sheets contain bounded integer metadata and
    are not rewritten. This does not evaluate formulas or apply display units.
    """
    with ZipFile(path) as archive:
        workbook = _xml(archive.read('xl/workbook.xml'))
        relationships = _xml(archive.read('xl/_rels/workbook.xml.rels'))
        by_id = {}
        for rel in relationships:
            identifier = rel.get('Id')
            if identifier in by_id:
                raise ValueError('정산표에 중복된 시트 연결 ID가 있습니다.')
            target = rel.get('Target', '')
            target = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
            by_id[identifier] = target
        physical = set()
        for node in workbook.iter():
            if _tag(node) != 'SHEET':
                continue
            name = node.get('name')
            if name not in book.sheetnames or name in ('_DSD_META', '_DSD_CELLS'):
                continue
            relationship = node.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
            target = by_id.get(relationship, '')
            if (not target.startswith('xl/worksheets/') or not target.endswith('.xml')
                    or target not in archive.namelist() or target in physical):
                raise ValueError('정산표의 시트 연결을 확인해 주세요.')
            physical.add(target)
            sheet = book[name]
            root = _xml(archive.read(target))
            coordinates = set()
            for cell in root.iter():
                if _tag(cell) != 'C':
                    continue
                address = cell.get('r', '').upper()
                if address in coordinates:
                    raise ValueError('정산표에 중복된 셀 좌표가 있습니다.')
                coordinates.add(address)
                if cell.get('t', 'n') != 'n':
                    continue
                values = [node for node in cell if _tag(node) == 'V']
                if len(values) > 1:
                    raise ValueError('정산표 셀에 중복 숫자가 있습니다.')
                if not values or values[0].text is None:
                    continue
                actual = sheet[address]
                if actual.data_type == 'n':
                    actual.value = _decimal_value(values[0].text)


def inspect_workbook_cells(workbook_path):
    """Read a bounded value-only workbook; formulas/links/macros are rejected."""
    source = Path(workbook_path)
    before = _digest(source, MAX_WORKBOOK_BYTES)
    book = _load_workbook(source)
    try:
        cells, sheets = [], []
        visited = 0
        # The guide is not a mapped input in our own extraction format. This
        # marker only controls the picker display; full provenance is verified
        # against a separately selected DSD before any value is written.
        derived_picker = ('_DSD_META' in book.sheetnames and
                          any(row == ('schema', 'account-normalizer-dsd-excel/1')
                              for row in book['_DSD_META'].iter_rows(min_row=2, max_row=8, max_col=2, values_only=True)))
        for sheet in book:
            if sheet.title in ('_DSD_META', '_DSD_CELLS') or (derived_picker and sheet.title in ('변환안내', '본문')):
                continue
            sheets.append(sheet.title)
            if sheet.max_row > 100000 or sheet.max_column > 80:
                raise ValueError('정산표 시트는 10만 행·80열 이내로 준비해 주세요.')
            for row in sheet:
                visited += len(row)
                if visited > 1000000:
                    raise ValueError('정산표 전체 셀 범위의 한도를 넘었습니다.')
                for cell in row:
                    if cell.value is None:
                        continue
                    _value(cell.value)
                    value = cell.value
                    if isinstance(value, Decimal):
                        value = int(value) if value == value.to_integral_value() else str(value)
                    cells.append({'sheet': sheet.title, 'cell': cell.coordinate,
                                  'value': value, 'type': 'number' if cell.data_type == 'n' else 'text',
                                  'numberFormat': cell.number_format})
                    if len(cells) > MAX_CELLS:
                        raise ValueError('정산표 값이 있는 셀은 10만 개 이하여야 합니다.')
        if _digest(source, MAX_WORKBOOK_BYTES) != before:
            raise ValueError('읽는 동안 정산표가 변경되었습니다. 다시 열어 주세요.')
        return {'schema': 'account-normalizer-workbook-cells/1', 'sourceHash': before,
                'sheets': sheets, 'cells': cells, 'cellCount': len(cells),
                'warnings': ['수식은 계산하지 않습니다. 확정값 정산표를 사용하고 표시 단위·백분율을 직접 확인해 주세요.']}
    finally:
        book.close()


def _scale(value):
    if not isinstance(value, str) or len(value) > 40:
        raise ValueError('표시 배율은 숫자 문자열이어야 합니다.')
    try:
        number = Decimal(value)
        if not number.is_finite() or not Decimal('0.000000000001') <= number <= Decimal('1000000000000'):
            raise ValueError('표시 배율은 10^-12~10^12 범위여야 합니다.')
        return number
    except InvalidOperation as exc:
        raise ValueError('표시 배율을 확인해 주세요.') from exc


def _validate_mappings(items, inventory, book):
    if not isinstance(items, list) or not items or len(items) > MAX_CELLS:
        raise ValueError('확인한 셀 연결이 1~10만 개 필요합니다.')
    actual = {cell['id']: cell for cell in inventory['cells']}
    seen, result = set(), []
    for item in items:
        if not isinstance(item, dict) or item.get('confirmed') is not True:
            raise ValueError('정산표와 DSD의 각 셀 연결을 직접 확인해 주세요.')
        target = item.get('dsdCellId')
        sheet, address = item.get('sheet'), item.get('cell')
        if not isinstance(target, str) or target not in actual or target in seen:
            raise ValueError('DSD 셀 연결이 없거나 중복됩니다.')
        if not isinstance(sheet, str) or sheet not in book.sheetnames or sheet in ('_DSD_META', '_DSD_CELLS'):
            raise ValueError('정산표 연결 시트를 확인해 주세요.')
        if not isinstance(address, str) or not _COORD.fullmatch(address):
            raise ValueError('정산표 연결 셀 좌표를 확인해 주세요.')
        row, col = coordinate_to_tuple(address)
        if row > 100000 or col > 80:
            raise ValueError('정산표 연결 셀이 지원 범위를 넘었습니다.')
        source = book[sheet][address]
        if source.__class__.__name__ == 'MergedCell':
            raise ValueError('병합 셀은 왼쪽 위의 실제 값 셀에 연결해 주세요.')
        _value(source.value)
        scale = str(_scale(item.get('scale', '1')))
        result.append({'sheet': sheet, 'cell': address, 'dsdCellId': target,
                       'confirmed': True, 'scale': scale})
        seen.add(target)
    return result


def build_workbook_mapping(template_path, workbook_path, mappings):
    """Bind explicitly confirmed coordinates to the current input bytes."""
    hashes = {'template': _digest(template_path), 'workbook': _digest(workbook_path, MAX_WORKBOOK_BYTES)}
    inventory = read_dsd(template_path)
    book = _load_workbook(workbook_path)
    try:
        result = _validate_mappings(mappings, inventory, book)
    finally:
        book.close()
    if hashes != {'template': _digest(template_path), 'workbook': _digest(workbook_path, MAX_WORKBOOK_BYTES)}:
        raise ValueError('셀 연결 확인 중 원본이 변경되었습니다. 다시 불러와 주세요.')
    return {'schema': MAPPING_SCHEMA, 'sourceHashes': hashes, 'mappings': result}


def _plain(number):
    text = format(number, 'f')
    return text.rstrip('0').rstrip('.') if '.' in text else text


def _percent_format(number_format, value):
    """Recognise an active percent operator, excluding quoted/escaped literals."""
    sections, current, quoted, escaped, bracketed = [], [], False, False, False
    for char in number_format:
        if escaped:
            escaped = False
            continue
        if char == '\\':
            escaped = True
            continue
        if char == '"':
            quoted = not quoted
            continue
        if quoted:
            continue
        if char == '[':
            bracketed = True
        if char == ']':
            bracketed = False
            continue
        if bracketed:
            if char in '<>=':
                # Conditional format selection needs an Excel format engine.
                # Do not guess a different display value in this converter.
                current.append('condition')
            continue
        if char == ';':
            sections.append(current); current = []
        elif char == '%':
            current.append('%')
    sections.append(current)
    active = sections[1] if value < 0 and len(sections) > 1 else sections[2] if value == 0 and len(sections) > 2 else sections[0]
    if any('condition' in section for section in sections) and any('%' in section for section in sections):
        raise ValueError('조건부 백분율 서식은 표시값을 확정 텍스트로 준비해 주세요.')
    if active.count('%') > 1:
        raise ValueError('중복 백분율 서식은 표시값을 확정 텍스트로 준비해 주세요.')
    return '%' in active


def _mapped_text(cell, scale, original=None):
    value = cell.value
    number = _scale(scale)
    if number != 1:
        if value is None or isinstance(value, (str, bool)):
            raise ValueError('숫자 셀에만 표시 배율을 적용할 수 있습니다.')
        with localcontext() as context:
            context.prec = 250
            return _plain(Decimal(_value(value)) * number)
    if original is not None:
        from xbrl_filing.dsd_excel import export_cell_value
        exported, kind, number_format = export_cell_value(original)
        numeric_equal = (isinstance(value, (int, float, Decimal)) and not isinstance(value, bool)
                         and isinstance(exported, (int, float, Decimal)) and not isinstance(exported, bool)
                         and Decimal(str(value)) == Decimal(str(exported)))
        if numeric_equal or value == exported or (value in (None, '') and exported in (None, '')):
            return original['text']
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = Decimal(_value(value))
        # A percent-format workbook stores a ratio. Preserve its explicit
        # display convention when no additional user scale is applied.
        if _percent_format(cell.number_format, number):
            with localcontext() as context:
                context.prec = 250
                return _plain(number * 100) + '%'
        return _plain(number)
    return _value(value)


def _derived_mappings(book, inventory, template_path, digest):
    from xbrl_filing.dsd_excel import (SCHEMA, META_SHEET, CELLS_SHEET, META_HEADERS,
                                      CELL_HEADERS, build_export_plan, cell_metadata_record)
    if META_SHEET not in book.sheetnames or CELLS_SHEET not in book.sheetnames:
        raise ValueError('일반 정산표는 확인한 셀 연결을 지정해 주세요. DSD에서 추출한 엑셀은 원본 템플릿을 선택해 주세요.')
    meta, records = book[META_SHEET], book[CELLS_SHEET]
    if tuple(cell.value for cell in meta[1]) != tuple(META_HEADERS) or tuple(cell.value for cell in records[1]) != tuple(CELL_HEADERS):
        raise ValueError('DSD 추출 엑셀의 메타데이터 머리글이 다릅니다.')
    if meta.max_row > 300 or meta.max_column != 2 or records.max_column != len(CELL_HEADERS):
        raise ValueError('DSD 추출 엑셀의 메타데이터 범위를 확인해 주세요.')
    values, tables = {}, []
    for key, value in meta.iter_rows(min_row=2, values_only=True):
        if key == 'table':
            if not isinstance(value, str) or len(value) > 10000:
                raise ValueError('DSD 표 메타데이터 형식이 다릅니다.')
            tables.append(json.loads(value))
        elif key in values:
            raise ValueError('DSD 추출 엑셀에 중복 메타데이터가 있습니다.')
        else:
            values[key] = value
    expected = {'schema': SCHEMA, 'sourceSha256': digest, 'sourceName': Path(template_path).name,
                'sourceBytes': Path(template_path).stat().st_size, 'tableCount': len(inventory['tables']),
                'cellCount': len(inventory['cells'])}
    plan = build_export_plan(inventory)
    if values != expected or tables != plan or records.max_row != len(inventory['cells']) + 1:
        raise ValueError('DSD 추출 엑셀의 원본·표·셀 메타데이터가 현재 템플릿과 다릅니다.')
    by_table = {table['id']: table for table in plan}
    grouped_cells = {table['id']: [] for table in plan}
    for original in inventory['cells']:
        grouped_cells[original['id'].split(':r', 1)[0]].append(original)
    mappings = []
    # Summary sheets carry no mapped inputs and are not trusted as metadata.
    for row, original in zip(records.iter_rows(min_row=2, values_only=True), inventory['cells']):
        position = original['source']
        table_id = original['id'].split(':r', 1)[0]
        sheet_name = by_table[table_id]['sheet']
        expected_row = cell_metadata_record(original, sheet_name)
        normalized = [value if value is not None else '' for value in row]
        if normalized != [value if value is not None else '' for value in expected_row]:
            raise ValueError('DSD 추출 엑셀의 셀 출처가 변경되었습니다. 다시 추출해 주세요.')
        sheet = book[sheet_name] if sheet_name in book.sheetnames else None
        if sheet is None:
            raise ValueError('DSD 추출 엑셀의 표 시트가 없습니다.')
        address = str(expected_row[2])
        mappings.append({'sheet': sheet_name, 'cell': address, 'dsdCellId': original['id'],
                         'confirmed': True, 'scale': '1'})
    for table in plan:
        sheet = book[table['sheet']]
        actual_merges = {str(value) for value in sheet.merged_cells.ranges}
        from openpyxl.utils import get_column_letter
        expected_merges = set()
        anchors = set()
        for original in grouped_cells[table['id']]:
            p = original['source']
            anchors.add((p['row'], p['column']))
            if p['rowspan'] > 1 or p['colspan'] > 1:
                start = f"{get_column_letter(p['column'])}{p['row']}"
                end = f"{get_column_letter(p['column']+p['colspan']-1)}{p['row']+p['rowspan']-1}"
                expected_merges.add(f'{start}:{end}')
        if actual_merges != expected_merges:
            raise ValueError('DSD 추출 엑셀의 병합 구조가 바뀌었습니다. 표 구조를 유지해 주세요.')
        if sheet.max_row > table['rows'] or sheet.max_column > table['columns']:
            raise ValueError('DSD 표 범위 밖에 값이나 서식이 추가되었습니다. 원본 구조를 유지해 주세요.')
        # Inspect only physical OOXML cells already loaded by openpyxl; using
        # iter_rows on a sparse 100k-by-80 range would instantiate empty cells.
        for position, cell in sheet._cells.items():
            if position not in anchors and cell.value not in (None, ''):
                raise ValueError('원본 DSD 셀이 없는 위치에 값이 추가되었습니다. 새 셀은 DART 편집기에서 작성해 주세요.')
    return mappings


def _nodes(root, member):
    """Same anchor ordering as the existing safe reader, without guessing tags."""
    result = {}
    table_number = 0
    for table in root.iter():
        if _tag(table) != 'TABLE':
            continue
        table_number += 1
        rows = [node for node in table.iter() if _tag(node) == 'TR'
                and next((p for p in node.iterancestors() if _tag(p) == 'TABLE'), None) is table]
        occupied = set()
        for row_number, row in enumerate(rows, 1):
            column = 1
            for node in row:
                if _tag(node) not in ('TD', 'TH', 'TE', 'TU'):
                    continue
                while (row_number, column) in occupied:
                    column += 1
                colspan = int(node.get('COLSPAN', node.get('colspan', '1')))
                rowspan = int(node.get('ROWSPAN', node.get('rowspan', '1')))
                result[f'{member}#t{table_number}:r{row_number}:c{column}'] = node
                occupied.update((r, c) for r in range(row_number, row_number + rowspan)
                                for c in range(column, column + colspan))
                column += colspan
    return result


def _patch_xml(data, member, changes):
    root = _xml(data)
    nodes = _nodes(root, member)
    for identifier, text in changes.items():
        node = nodes.get(identifier)
        if node is None:
            raise ValueError('DSD 템플릿에서 연결한 셀을 찾을 수 없습니다.')
        value_node = node
        while len(value_node):
            if (len(value_node) != 1 or (value_node.text or '').strip()
                    or (value_node[0].tail or '').strip()
                    or _tag(value_node[0]) not in ('P', 'SPAN', 'FONT', 'B', 'I', 'U', 'STRONG', 'EM', 'SUP', 'SUB')):
                raise ValueError('여러 서식·중첩 태그에 텍스트가 나뉜 셀은 자동 수정하지 않습니다. DART 편집기에서 수정해 주세요.')
            value_node = value_node[0]
        value_node.text = text or None
    info = root.getroottree().docinfo
    encoding = info.encoding or 'UTF-8'
    # Preserve processing instructions outside DOCUMENT as well as unknown
    # elements/attributes. Only member(s) containing changed values are written.
    declaration = bool(re.match(br'\s*<\?xml\s', data)) or data.startswith((b'\xff\xfe', b'\xfe\xff'))
    output = etree.tostring(root.getroottree(), encoding=encoding, xml_declaration=declaration)
    if len(output) > MAX_XML_BYTES:
        raise ValueError('수정된 DSD XML의 크기 한도를 넘었습니다.')
    return output


def _verify_result(original, output_path, changes):
    after = read_dsd(output_path)
    if len(after['tables']) != len(original['tables']) or len(after['cells']) != len(original['cells']):
        raise ValueError('DSD 초안의 표·셀 개수가 원본과 다릅니다.')
    for before, after_cell in zip(original['cells'], after['cells']):
        if before['source'] != after_cell['source']:
            # Line numbers can legitimately shift for newly inserted newlines;
            # anchor coordinates and spans must remain exactly the same.
            a, b = dict(before['source']), dict(after_cell['source'])
            a.pop('xmlLine', None); b.pop('xmlLine', None)
            if a != b:
                raise ValueError('DSD 초안의 셀 위치·병합 구조가 변경되었습니다.')
        if after_cell['text'] != changes.get(before['id'], before['text']).strip():
            raise ValueError('DSD 초안의 셀 값이 확인한 정산표 값과 다릅니다.')
    return after


def export_workbook_to_dsd(template_path, workbook_path, output_path, mappings=None):
    """Apply explicitly mapped values to an existing document, then reopen it.

    General workbooks require a hash-bound envelope produced by
    ``build_workbook_mapping``. A DSD-derived workbook may use its verified
    provenance for full-grid mapping. No tags, periods, units or company data
    are inferred, and no claim is made about official editor compatibility.
    """
    template, workbook, target = Path(template_path), Path(workbook_path), Path(output_path)
    if target.exists():
        raise FileExistsError('기존 DSD 파일을 덮어쓸 수 없습니다. 새 파일 이름을 선택해 주세요.')
    if target.resolve() in (template.resolve(), workbook.resolve()):
        raise ValueError('원본 파일을 덮어쓸 수 없습니다.')
    if target.suffix.casefold() != template.suffix.casefold():
        raise ValueError('원본 템플릿과 같은 확장자(.dsd/.xml/.zip)로 저장해 주세요.')
    hashes = {'template': _digest(template), 'workbook': _digest(workbook, MAX_WORKBOOK_BYTES)}
    inventory = read_dsd(template)
    book = _load_workbook(workbook)
    derived = mappings is None
    try:
        if derived:
            selected = _derived_mappings(book, inventory, template, hashes['template'])
        elif (not isinstance(mappings, dict) or mappings.get('schema') != MAPPING_SCHEMA
              or mappings.get('sourceHashes') != hashes):
            raise ValueError('정산표·DSD 원본 해시가 일치하는 확인된 셀 연결이 필요합니다.')
        else:
            selected = _validate_mappings(mappings.get('mappings'), inventory, book)
        by_id = {cell['id']: cell for cell in inventory['cells']}
        changes = {}
        for mapping in selected:
            original = by_id[mapping['dsdCellId']]
            text = _mapped_text(book[mapping['sheet']][mapping['cell']], mapping['scale'], original if derived else None)
            if text != original['text']:
                changes[original['id']] = text
    finally:
        book.close()
    target.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(prefix='.dsd-draft-', suffix=target.suffix, dir=target.parent)
    os.close(handle)
    temp = Path(temporary)
    untouched_members, changed_members = 0, 0
    changes_by_member = {}
    for identifier, text in changes.items():
        member = by_id[identifier]['source']['member'].casefold()
        changes_by_member.setdefault(member, {})[identifier] = text
    try:
        with template.open('rb') as source:
            signature = source.read(4)
        if signature in (b'PK\x03\x04', b'PK\x05\x06'):
            with ZipFile(template) as archive, ZipFile(temp, 'w') as output:
                index, _ = _inventory(archive, _Budget())
                output.comment = archive.comment
                for entry in archive.infolist():
                    member = entry.filename
                    member_changes = changes_by_member.get(member.casefold(), {})
                    if member_changes:
                        data = _read_member(archive, index, member, MAX_XML_BYTES)
                        data = _patch_xml(data, member.casefold(), member_changes)
                        changed_members += 1
                        output.writestr(copy(entry), data)
                    else:
                        untouched_members += 1
                        if entry.is_dir():
                            output.writestr(copy(entry), b'')
                        else:
                            with archive.open(entry) as incoming, output.open(copy(entry), 'w') as outgoing:
                                shutil.copyfileobj(incoming, outgoing, length=1024 * 1024)
            with ZipFile(template) as before, ZipFile(temp) as after:
                if before.namelist() != after.namelist():
                    raise ValueError('DSD 압축 구성요소의 순서·이름이 변경되었습니다.')
                for entry in before.infolist():
                    if entry.filename.casefold() not in changes_by_member:
                        with before.open(entry) as left, after.open(entry.filename) as right:
                            before_hash = hashlib.file_digest(left, 'sha256').digest()
                            after_hash = hashlib.file_digest(right, 'sha256').digest()
                        if before_hash != after_hash:
                            raise ValueError('변경하지 않은 DSD 구성요소가 원본과 다릅니다.')
        else:
            data = template.read_bytes()
            temp.write_bytes(_patch_xml(data, template.name, changes) if changes else data)
            changed_members = int(bool(changes)); untouched_members = int(not changes)
            # A standalone XML member ID is its source basename. Reopening the
            # new file changes only this identity; adapt the verification copy.
            inventory = {**inventory, 'cells': [{**cell, 'source': {**cell['source'], 'member': temp.name}}
                                               for cell in inventory['cells']]}
        if signature not in (b'PK\x03\x04', b'PK\x05\x06'):
            check = read_dsd(temp)
            if len(check['cells']) != len(inventory['cells']):
                raise ValueError('DSD XML 초안의 셀 개수가 다릅니다.')
            for original, result in zip(inventory['cells'], check['cells']):
                wanted = changes.get(original['id'], original['text']).strip()
                if result['text'] != wanted or any(original['source'][k] != result['source'][k] for k in ('table','row','column','rowspan','colspan')):
                    raise ValueError('DSD XML 초안의 값·구조가 확인한 내용과 다릅니다.')
        else:
            _verify_result(inventory, temp, changes)
        if hashes != {'template': _digest(template), 'workbook': _digest(workbook, MAX_WORKBOOK_BYTES)}:
            raise ValueError('DSD 초안 작성 중 원본이 변경되었습니다. 다시 불러와 주세요.')
        os.link(temp, target)
    finally:
        temp.unlink(missing_ok=True)
    return {'schema': 'account-normalizer-dsd-workbook-output/1', 'output': str(target),
            'sourceHashes': hashes, 'sha256': _digest(target),
            'mappedCellCount': len(selected), 'changedCellCount': len(changes),
            'unchangedCellCount': len(inventory['cells']) - len(changes),
            'unmappedCellCount': len(inventory['cells']) - len(selected),
            'changedMemberCount': changed_members, 'untouchedMemberCount': untouched_members,
            'sourceUnchanged': True, 'reopenedAndChecked': True,
            'officialEditorOpened': False, 'officialValidated': False, 'submissionApproved': False,
            'warnings': ['기존 DSD/XML 템플릿의 확인한 표 셀 값만 바꾼 초안입니다. 엑셀의 변환안내·본문 참고 시트 편집은 자동 반영하지 않습니다. 회사·기간·표시 단위·주석 연결과 변경하지 않은 항목을 검토해 주세요.',
                         'DART 편집기에서 열기·파일형식검사를 수행하지 않았습니다. 전송파일 생성·제출 승인·접수를 뜻하지 않습니다.']}

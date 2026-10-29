"""DSD XML table export for the desktop application; no Excel process needed."""
import json
import os
from pathlib import Path
import tempfile
from zipfile import BadZipFile

from lxml import etree
from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from io_xlsx.package_guard import check_package
from ..disclosure_compare.dsd import read_dsd
from .plan import (
    SCHEMA, META_SHEET, CELLS_SHEET, META_HEADERS, CELL_HEADERS,
    build_export_plan, cell_metadata_record, export_cell_value, source_digest,
)
from .plan import _text_fits


def _write(cell, value):
    cell.value = value
    if isinstance(value, str):
        # Formula-looking source descriptions remain OOXML inline/shared text.
        cell.data_type = 's'


def _row(sheet, values):
    # Worksheet.max_row scans every existing cell. Using append avoids an
    # accidental quadratic scan over the 100,000-row provenance worksheet.
    values_to_append = []
    for column, value in enumerate(values, 1):
        cell = Cell(sheet, row=1, column=column)
        _write(cell, value)
        values_to_append.append(cell)
    sheet.append(values_to_append)


def _header(sheet, row=1):
    for cell in sheet[row]:
        cell.font = Font(name='맑은 고딕', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='203B5A')
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[row].height = 30


def _hint_text(values):
    text = ' / '.join(values)
    encoded = text.encode('utf-16-le')
    if len(encoded) > 64000:
        return encoded[:63996].decode('utf-16-le', errors='ignore') + ' …'
    return text


def _narrative_sheet(workbook, narratives):
    sheet = workbook.create_sheet('본문')
    sheet.sheet_view.showGridLines = False
    _row(sheet, ['원문 참고용 본문이에요. 이 시트의 편집은 DSD에 자동 반영하지 않아요. 이미지·첨부와 시각적 페이지 배치는 원본에서 확인해 주세요.'])
    sheet.merge_cells('A1:E1')
    sheet['A1'].font = Font(name='맑은 고딕', size=11, bold=True, color='203B5A')
    sheet['A1'].fill = PatternFill('solid', fgColor='EAF1FF')
    sheet['A1'].alignment = Alignment(vertical='center', wrap_text=True)
    sheet.row_dimensions[1].height = 44
    _row(sheet, ['순서', '문서 XML 멤버', '절제목', '원문문구', '원본위치'])
    _header(sheet, 2)
    for order, item in enumerate(narratives, 1):
        location = f'{item["id"]} · {item["tag"]} · XML {item.get("xmlLine") or "?"}행'
        _row(sheet, [order, item['member'], ' / '.join(item.get('sectionPathFull', item.get('sectionPath', []))),
                     item['text'], location])
        row_number = order + 2
        for column in range(1, 6):
            cell = sheet.cell(row_number, column)
            cell.font = Font(name='맑은 고딕', size=11, color='172B4D')
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        sheet.row_dimensions[row_number].height = min(180, max(32, 22 * (1 + len(item['text']) // 90 + item['text'].count('\n'))))
    for column, width in [('A', 10), ('B', 38), ('C', 44), ('D', 105), ('E', 48)]:
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = 'D3'
    if narratives:
        sheet.auto_filter.ref = f'A2:E{len(narratives) + 2}'
    else:
        _row(sheet, ['', '', '', '읽을 수 있는 표 밖 XML 텍스트가 없어요.', ''])
    return sheet


def _commit(temp_path, target):
    # os.rename is exclusive on Windows; hard linking is exclusive on POSIX.
    # No fallback ever replaces an existing file.
    if os.name == 'nt':
        os.rename(temp_path, target)
    else:
        os.link(temp_path, target)
        temp_path.unlink()


def export_dsd_excel(source_path, output_path):
    """Export all supported tables and source-bound provenance to a new XLSX.

    Unit and period strings are kept as hints; no unit conversion, amount
    aggregation, approval, taxonomy selection or prior-year shift is inferred.
    The caller can edit existing anchor cells and use the original DSD as a
    template through the separate Excel-to-DSD workflow.
    """
    source, target = Path(source_path), Path(output_path)
    if target.suffix.lower() != '.xlsx':
        raise ValueError('저장 파일은 .xlsx 확장자를 사용해 주세요.')
    if source.resolve() == target.resolve():
        raise ValueError('원본 파일에 덮어쓸 수 없습니다.')
    if target.exists():
        raise FileExistsError('기존 Excel 파일에 덮어쓸 수 없습니다. 새 이름으로 저장해 주세요.')
    before = source_digest(source)
    source_bytes = source.stat().st_size
    try:
        inventory = read_dsd(source)
    except (etree.XMLSyntaxError, BadZipFile) as error:
        raise ValueError('올바른 XML 표를 포함한 DSD/공시 원문을 선택해 주세요.') from error
    plan = build_export_plan(inventory)
    # The document reader shares the same bounded XML/ZIP parser used by DSD
    # version comparison; no network or Excel process is invoked.
    from ..dsd_document import extract_dsd_narratives
    narrative_report = extract_dsd_narratives(source)
    narratives = narrative_report['narratives']
    for item in narratives:
        for text in (item['member'], item['text'],
                     ' / '.join(item.get('sectionPathFull', item.get('sectionPath', []))),
                     f'{item["id"]} · {item["tag"]} · XML {item.get("xmlLine") or "?"}행'):
            _text_fits(text)
    if source_digest(source) != before:
        raise ValueError('읽는 동안 DSD 원본이 바뀌었습니다. 다시 선택해 주세요.')
    workbook = Workbook()
    guide = workbook.active
    guide.title = '변환안내'
    guide.sheet_view.showGridLines = False
    _row(guide, ['DSD 표를 Excel로 가져왔어요', ''])
    _row(guide, ['원본', source.name])
    _row(guide, ['원본 SHA-256', before])
    _row(guide, ['자료 범위', 'XML 표의 텍스트·숫자·셀 병합과 표 밖 본문을 가져왔어요. 본문 시트는 참고용이며 문서 서식·이미지·페이지 배치는 원본 DSD에서 확인해 주세요.'])
    _row(guide, ['값 편집', '기존 셀의 값만 편집해 주세요. 행·열·표 추가/삭제와 셀 병합 변경은 DSD에 반영하지 않아요.'])
    _row(guide, ['숫자와 단위', '원본 표시 단위를 유지해요. 공란·대시·0은 다르며 전기 금액을 당기로 옮기거나 자동으로 합산하지 않아요.'])
    _row(guide, ['DSD로 반영', '별도 Excel→DSD 기능에서 이 Excel과 같은 원본 DSD를 선택해 주세요. 숨긴 대응 정보도 원본과 다시 검사해요. 변환안내·본문 시트의 편집은 자동 반영하지 않아요.'])
    _row(guide, ['공식 검토', '생성 파일은 편집용 초안이에요. DART 편집기에서 서식·문맥·수치를 확인하고 공식 검사를 진행해 주세요.'])
    _row(guide, ['정밀도', 'Excel의 15자리 숫자 정밀도를 넘는 값과 선행 0 표기는 텍스트로 보존해요. 수식은 생성하지 않아요.'])
    _row(guide, [])
    _row(guide, ['표 시트', 'XML 멤버', '표 번호', '행', '열', '원본 셀', '단위 안내', '기간 안내'])
    _header(guide, 1)
    _header(guide, 11)
    guide.column_dimensions['A'].width = 34
    guide.column_dimensions['B'].width = 78
    for column in ('C', 'D', 'E', 'F'):
        guide.column_dimensions[column].width = 13
    guide.column_dimensions['G'].width = 40
    guide.column_dimensions['H'].width = 45
    for row in guide.iter_rows(min_row=2, max_row=9):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.font = Font(name='맑은 고딕', size=11, color='172B4D')
        guide.row_dimensions[row[0].row].height = 44
    guide.freeze_panes = 'B12'
    _narrative_sheet(workbook, narratives)
    meta = workbook.create_sheet(META_SHEET)
    _row(meta, META_HEADERS)
    for key, value in [('schema', SCHEMA), ('sourceSha256', before),
                       ('sourceName', source.name), ('sourceBytes', source_bytes),
                       ('tableCount', len(plan)), ('cellCount', len(inventory['cells']))]:
        _row(meta, [key, value])
    cells_sheet = workbook.create_sheet(CELLS_SHEET)
    _row(cells_sheet, CELL_HEADERS)
    meta.sheet_state = cells_sheet.sheet_state = 'veryHidden'
    by_id = {cell['id']: cell for cell in inventory['cells']}
    precision_count, precision_examples = 0, []
    border = Border(left=Side(style='hair', color='DCE3ED'),
                    right=Side(style='hair', color='DCE3ED'),
                    top=Side(style='hair', color='DCE3ED'),
                    bottom=Side(style='hair', color='DCE3ED'))
    for table, descriptor in zip(inventory['tables'], plan):
        _row(meta, ['table', json.dumps(descriptor, ensure_ascii=False, sort_keys=True, separators=(',', ':'))])
        _row(guide, [descriptor['sheet'], descriptor['member'], descriptor['number'],
                     descriptor['rows'], descriptor['columns'], descriptor['cellCount'],
                     _hint_text(table.get('unitHints', [])),
                     _hint_text(table.get('periodHints', []))])
        sheet = workbook.create_sheet(descriptor['sheet'])
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 90
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.outlinePr.summaryRight = False
        sheet.freeze_panes = 'B2'
        for column in range(1, descriptor['columns'] + 1):
            sheet.column_dimensions[get_column_letter(column)].width = 32 if column == 1 else 21
        for cell_id in table['cellIds']:
            original = by_id[cell_id]
            position = original['source']
            cell = sheet.cell(position['row'], position['column'])
            value, kind, number_format = export_cell_value(original)
            _write(cell, value)
            cell.number_format = number_format
            cell.font = Font(name='맑은 고딕', size=11, color='172B4D')
            cell.alignment = Alignment(vertical='center', wrap_text=True,
                                       horizontal='right' if kind == 'number' else 'left')
            cell.border = border
            if original.get('isHeader'):
                cell.fill = PatternFill('solid', fgColor='EAF1FF')
                cell.font = Font(name='맑은 고딕', size=11, bold=True, color='203B5A')
            sheet.row_dimensions[position['row']].height = max(
                sheet.row_dimensions[position['row']].height or 24,
                min(120, 24 * (1 + original['text'].count('\n'))))
            if position['rowspan'] > 1 or position['colspan'] > 1:
                sheet.merge_cells(start_row=position['row'], start_column=position['column'],
                                  end_row=position['row'] + position['rowspan'] - 1,
                                  end_column=position['column'] + position['colspan'] - 1)
            _row(cells_sheet, cell_metadata_record(original, descriptor['sheet']))
            if kind == 'precision-text':
                precision_count += 1
                if len(precision_examples) < 20:
                    precision_examples.append(cell_id)
        if descriptor['rows'] and descriptor['columns']:
            sheet.print_area = f'A1:{get_column_letter(descriptor["columns"])}{descriptor["rows"]}'
    target.parent.mkdir(parents=True, exist_ok=True)
    handle, name = tempfile.mkstemp(prefix='.dsd-excel-', suffix='.xlsx', dir=target.parent)
    os.close(handle)
    temporary = Path(name)
    try:
        workbook.save(temporary)
        workbook.close()
        check_package(temporary)
        if source_digest(source) != before:
            raise ValueError('저장하는 동안 DSD 원본이 바뀌었습니다. 결과를 저장하지 않았어요.')
        output_hash = source_digest(temporary)
        _commit(temporary, target)
    finally:
        workbook.close()
        temporary.unlink(missing_ok=True)
    warnings = list(inventory.get('warnings', [])) + list(narrative_report.get('warnings', []))
    warnings.append('본문은 XML의 표 밖 텍스트 참고용이에요. 이미지·첨부·시각적 페이지 배치와 표 안 셀 밖 서술은 보존하지 않으며, 본문 시트의 편집은 DSD에 자동 반영하지 않아요.')
    if precision_count:
        warnings.append(f'{precision_count}개 숫자는 Excel의 정밀도 한도를 넘거나 안전하게 표현할 수 없어 원문 텍스트로 보존했어요.')
    return {'schema': SCHEMA, 'operation': 'dsd-to-excel', 'status': 'exported',
            'sourceName': source.name, 'sourceSha256': before, 'sourceUnchanged': True,
            'output': str(target), 'outputSha256': output_hash,
            'counts': {'tables': len(plan), 'cells': len(inventory['cells']), 'narratives': len(narratives),
                       'precisionTextCells': precision_count},
            'precisionTextExamples': precision_examples, 'tables': plan,
            'warnings': warnings, 'officialGenerated': False,
            'validationPassed': False, 'finalSubmission': False}

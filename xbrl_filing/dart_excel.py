"""Read DART's published Excel concept catalogue, without inventing an XSD.

The 2026 workbook has prefix/local-name columns but no namespace URI column.
Consequently namespace URIs remain unresolved until a real DTS is loaded.
Neither a workbook title nor this parser authenticates an official taxonomy.
"""
import hashlib
from io import BytesIO
from pathlib import Path, PurePosixPath
import re
from xml.parsers import expat
from zipfile import BadZipFile, ZipFile

from core.errors import InputError

MAX_BYTES = 50 * 1024 * 1024
MAX_EXPANDED = 256 * 1024 * 1024
MAX_ENTRY_BYTES = 128 * 1024 * 1024
MAX_ENTRIES = 1000
MAX_CONCEPTS = 20000
MAX_ROLES = 10000
MAX_PRESENTATION_ROWS = 100000
MAX_RELATIONSHIPS = 100000
MAX_TEXT = 20000
_NAME = re.compile(r'[A-Za-z_][A-Za-z0-9_.-]*\Z')
_NUMERIC = {'monetaryItemType', 'decimalItemType', 'floatItemType',
            'doubleItemType', 'integerItemType', 'nonPositiveIntegerItemType',
            'negativeIntegerItemType', 'longItemType', 'intItemType',
            'shortItemType', 'byteItemType', 'nonNegativeIntegerItemType',
            'unsignedLongItemType', 'unsignedIntItemType', 'unsignedShortItemType',
            'unsignedByteItemType', 'positiveIntegerItemType', 'pureItemType',
            'sharesItemType', 'perShareItemType', 'percentItemType'}


def _package(data):
    """Inspect the whole archive before openpyxl parses any XML."""
    try:
        with ZipFile(BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_ENTRIES or sum(e.file_size for e in entries) > MAX_EXPANDED:
                raise InputError('택사노미 엑셀의 압축해제 크기·파일 수 한도를 넘었습니다.')
            names = set()
            for entry in entries:
                name = entry.filename
                path = PurePosixPath(name)
                if (name.casefold() in names or '\\' in entry.orig_filename
                        or any(ord(char) < 32 for char in entry.orig_filename) or path.is_absolute()
                        or '..' in path.parts or ':' in name
                        or entry.file_size > MAX_ENTRY_BYTES or entry.flag_bits & 1
                        or (entry.external_attr >> 16) & 0o170000 == 0o120000):
                    raise InputError('안전하지 않은 택사노미 엑셀 압축 구조입니다.')
                names.add(name.casefold())
                lower = name.lower()
                if any(token in lower for token in ('externallinks/', 'vba', 'embeddings/', 'connections.xml')):
                    raise InputError('매크로·외부 연결·내장 개체가 있는 택사노미 엑셀은 지원하지 않습니다.')
                if not lower.endswith(('.xml', '.rels')):
                    continue
                parser = expat.ParserCreate(namespace_separator='}')
                def forbidden(*unused):
                    raise InputError('DOCTYPE·XML 엔터티가 있는 택사노미 엑셀은 지원하지 않습니다.')
                parser.StartDoctypeDeclHandler = forbidden
                parser.EntityDeclHandler = forbidden
                parser.ExternalEntityRefHandler = forbidden
                if lower.startswith('xl/worksheets/'):
                    def start(tag, attrs):
                        local = tag.rsplit('}', 1)[-1]
                        if local == 'row':
                            value = attrs.get('r', '')
                            if not value.isascii() or not value.isdigit() or not 1 <= int(value) <= MAX_PRESENTATION_ROWS:
                                raise InputError('택사노미 시트는 100,000행 이내여야 합니다.')
                        elif local == 'c':
                            coordinate = attrs.get('r', '')
                            match = re.fullmatch(r'([A-Z]+)[1-9][0-9]*', coordinate)
                            if match is None:
                                raise InputError('택사노미 셀 좌표가 올바르지 않습니다.')
                            column = 0
                            for char in match[1]:column = column * 26 + ord(char) - 64
                            if column > 80:
                                raise InputError('택사노미 시트는 80열 이내여야 합니다.')
                    parser.StartElementHandler = start
                with archive.open(entry) as stream:
                    while block := stream.read(65536):parser.Parse(block, False)
                parser.Parse(b'', True)
    except (BadZipFile, expat.ExpatError, RuntimeError, KeyError) as exc:
        raise InputError('정상적인 택사노미 .xlsx 파일이 아닙니다.') from exc


def _text(cell):
    if cell.data_type in ('f', 'e'):
        raise InputError('택사노미 정보는 수식·오류 값 대신 확정값이어야 합니다.')
    value = cell.value
    value = '' if value is None else str(value).strip()
    if len(value) > MAX_TEXT:
        raise InputError('택사노미 셀의 문자 수 한도를 넘었습니다.')
    return value


def _rows(sheet, columns, limit):
    count = 0
    for row in sheet.iter_rows(max_col=columns):
        count += 1
        if count > limit:
            raise InputError(f'{sheet.title} 시트의 행 수 한도를 넘었습니다.')
        yield [_text(cell) for cell in row]


def _boolean(value):
    if value in ('', 'false', 'False', '0'):return False
    if value in ('true', 'True', '1'):return True
    raise InputError('택사노미의 abstract·nillable 값이 올바르지 않습니다.')


def _qname(prefix, name):
    if _NAME.fullmatch(prefix) is None or _NAME.fullmatch(name) is None:
        raise InputError('택사노미의 prefix·name이 XML 이름 형식에 맞지 않습니다.')
    return prefix + ':' + name


def _kind(type_name, group):
    local = type_name.rsplit(':', 1)[-1]
    if group.endswith(':dimensionItem'):return 'dimension', False
    if group.endswith(':hypercubeItem'):return 'hypercube', False
    if local in _NUMERIC:
        if local == 'monetaryItemType':kind = 'monetary'
        elif local == 'sharesItemType':kind = 'shares'
        elif local == 'perShareItemType':kind = 'perShare'
        elif local == 'percentItemType':kind = 'percent'
        elif 'Integer' in local or local in ('integerItemType', 'intItemType', 'longItemType', 'shortItemType', 'byteItemType'):kind = 'integer'
        else:kind = 'decimal'
        return kind, True
    if local in ('dateItemType', 'dateTimeItemType'):return 'date', False
    if local == 'booleanItemType':return 'boolean', False
    if local in ('enumerationItemType', 'enumerationSetItemType'):return 'enumeration', False
    if local in ('stringItemType', 'textBlockItemType', 'normalizedStringItemType', 'domainItemType'):return 'text', False
    return 'other', False


def _concepts(workbook):
    rows = iter(_rows(workbook['Concepts'], 16, MAX_CONCEPTS + 1))
    header = next(rows, [])
    required = ('prefix', 'name', 'id', 'type', 'substitutionGroup', 'balance', 'periodType', 'abstract', 'nillable')
    if any(header.count(key) != 1 for key in required):
        raise InputError('DART Concepts 시트의 필수 열을 찾을 수 없습니다.')
    indexes = {key:header.index(key) for key in required}
    result = {};excluded = {}
    for row in rows:
        if not any(row):continue
        value = {key:row[index] for key,index in indexes.items()}
        qname = _qname(value['prefix'], value['name'])
        if value['periodType'] not in ('', 'instant', 'duration') or value['balance'] not in ('', 'debit', 'credit'):
            raise InputError('택사노미 개념의 기간·차대 속성이 올바르지 않습니다.')
        kind,numeric = _kind(value['type'], value['substitutionGroup'])
        concept = {'label':value['name'], 'namespace':'', 'localName':value['name'],
                         'periodType':value['periodType'], 'abstract':_boolean(value['abstract']),
                         'nillable':_boolean(value['nillable']), 'type':value['type'],
                         'numeric':numeric, 'kind':kind, 'substitutionGroup':value['substitutionGroup'],
                         'balance':value['balance'], 'deprecated':False, 'id':value['id']}
        if qname in excluded:continue
        if qname in result and result[qname] != concept:
            del result[qname]
            excluded[qname] = {'qname':qname, 'reason':'conflicting-declaration'}
        else:result[qname] = concept
    if not result:raise InputError('택사노미 Concepts 시트에 개념이 없습니다.')
    return result,list(excluded.values())


def _labels(workbook, concepts):
    if 'Label Link' not in workbook.sheetnames:return
    rows = iter(_rows(workbook['Label Link'], 6, MAX_CONCEPTS + 10))
    header = None
    for unused in range(10):
        candidate = next(rows, [])
        if 'prefix' in candidate and 'name' in candidate and 'label' in candidate:
            header = candidate;break
    if header is None:raise InputError('Label Link 시트의 열을 찾을 수 없습니다.')
    indexes = {key:header.index(key) for key in ('prefix', 'name', 'label')}
    dep = header.index('deprecatedDateLabel') if 'deprecatedDateLabel' in header else None
    seen = set()
    for row in rows:
        if not any(row):continue
        if not row[indexes['prefix']] and not row[indexes['name']] and not row[indexes['label']]:continue
        qname = _qname(row[indexes['prefix']], row[indexes['name']])
        if qname not in concepts:continue
        label = row[indexes['label']]
        if qname in seen and label != concepts[qname]['label']:
            raise InputError('택사노미 표준 라벨이 중복되어 있습니다: ' + qname)
        seen.add(qname)
        if label:concepts[qname]['label'] = label
        concepts[qname]['deprecated'] = bool((dep is not None and row[dep]) or '(사용중단)' in label)


def _roles(workbook, concepts):
    result = {}
    if 'RoleTypes' not in workbook.sheetnames:return result
    rows = iter(_rows(workbook['RoleTypes'], 7, MAX_ROLES + 1))
    header = next(rows, [])
    required = ('id', 'roleURI', 'definition', 'Defined Schema Systemid')
    if any(header.count(key) != 1 for key in required):
        raise InputError('RoleTypes 시트의 필수 열을 찾을 수 없습니다.')
    indexes = {key:header.index(key) for key in required}
    for row in rows:
        if not any(row):continue
        value = {key:row[index] for key,index in indexes.items()};uri = value['roleURI']
        if not uri.startswith(('https://', 'http://')) or uri in result:
            raise InputError('택사노미 역할 URI가 올바르지 않거나 중복되어 있습니다.')
        result[uri] = {'id':value['id'], 'definition':value['definition'],
                       'schema':value['Defined Schema Systemid'], 'industry':'', 'scope':'',
                       'category':'', 'concepts':[]}
    if 'Role 일람표' in workbook.sheetnames:
        categories = {}
        for row in _rows(workbook['Role 일람표'], 6, MAX_ROLES + 2):
            if row[1] in ('', '업종'):continue
            match = re.match(r'\[([^\]]+)\]', row[5])
            if match:categories[match[1]] = {'industry':row[1], 'category':row[2], 'scope':row[3]}
        for role in result.values():
            match = re.match(r'\[([^\]]+)\]', role['definition'])
            if match:role.update(categories.get(match[1], {}))
    if 'Presentation Link' in workbook.sheetnames:
        current = None;seen = {};relationships = 0
        for row in _rows(workbook['Presentation Link'], 3, MAX_PRESENTATION_ROWS):
            if row[0] == 'LinkRole':
                current = row[1] if row[1] in result else None
                continue
            if current is None or row[0] in ('', 'Definition', 'prefix'):continue
            qname = _qname(row[0], row[1])
            if qname not in concepts:continue
            role_seen = seen.setdefault(current, set())
            if qname in role_seen:continue
            relationships += 1
            if relationships > MAX_RELATIONSHIPS:raise InputError('택사노미 역할·개념 연결 수 한도를 넘었습니다.')
            role_seen.add(qname);result[current]['concepts'].append(qname)
    return result


def read_dart_taxonomy(path):
    """Return display/search metadata; this never validates filing XML or resolves a DTS."""
    from openpyxl import load_workbook
    path = Path(path)
    if path.suffix.lower() != '.xlsx':raise InputError('DART 택사노미 .xlsx 파일을 선택하세요.')
    try:
        with path.open('rb') as stream:data = stream.read(MAX_BYTES + 1)
        if len(data) > MAX_BYTES:raise InputError('50MB 이하 DART 택사노미 .xlsx 파일을 선택하세요.')
        _package(data)
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        try:
            if 'Concepts' not in workbook.sheetnames:raise InputError('DART 택사노미 Concepts 시트가 없습니다.')
            if len(workbook.sheetnames) > 32:raise InputError('택사노미 시트 수 한도를 넘었습니다.')
            concepts,excluded = _concepts(workbook);_labels(workbook, concepts)
            roles = _roles(workbook, concepts)
        finally:workbook.close()
    except InputError:raise
    except (OSError, ValueError, KeyError, BadZipFile) as exc:
        raise InputError('DART 택사노미 엑셀을 읽을 수 없습니다.') from exc
    return {'namespaces':{}, 'concepts':concepts, 'roles':roles,
            'source':{'name':path.name, 'sha256':hashlib.sha256(data).hexdigest()},
            'validationScope':'concept-metadata-only', 'namespaceResolutionRequired':True,
            'excludedConcepts':excluded}

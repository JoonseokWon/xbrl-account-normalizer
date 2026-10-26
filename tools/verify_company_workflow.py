"""Reproduce compound journals, portable work, and a large CSV review export."""
import argparse,csv,json,sys,time
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT));sys.path.insert(0,str(ROOT/'tests'))
from workflow_fixtures import profile,session,support_config
from application.loading import load_batch
from application.configuration import default_catalog
from application.reviews import review_rows,apply_review
from application.exporting import export_session
from application.support_preparation import prepare_support
from application.audit_export import export_audit_input,prepare_audit_tables
from application.workfiles import save_work,load_work
from core.errors import InputError

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--out',required=True);parser.add_argument('--large-rows',type=int,default=100004)
    args=parser.parse_args();folder=Path(args.out).resolve();folder.mkdir(parents=True,exist_ok=False)
    small=folder/'compound';small.mkdir();work=session(small,True);config=support_config(small)
    work.support_bundle=prepare_support(config);work.support_setup=config;work.support_config=config['tables']
    saved=save_work(work,small/'portable.acwork','시험 검토자');restored=load_work(saved)
    try:
        target=export_audit_input(restored,restored.support_bundle,small/'audit')
        assert prepare_audit_tables(restored,restored.support_bundle)[2]['debitTotal']==150
        portable={'pendingRows':restored.result['counts']['pendingRows'],'supportSources':len(restored.support_bundle.source['sources']),'auditInput':(target/'통합감사_원장입력.xlsx').relative_to(folder).as_posix()}
    finally:restored._restored_temp.cleanup()
    rows=args.large_rows
    if rows<10002 or rows%4:raise ValueError('--large-rows는 10,004 이상이며 4의 배수여야 합니다.')
    files=[];large=folder/'large';large.mkdir()
    for part in range(2):
        path=large/f'month-{part+1}.csv';files.append(path)
        with path.open('w',encoding='utf-8-sig',newline='') as stream:
            writer=csv.writer(stream);writer.writerow(['Corp','Day','Voucher','Seq','Code','Account','DR','CR','Bank','AccountNo','Party','Dept','Project','Who'])
            for i in range(rows//4):
                journal=f'M{part+1}-{i:06d}';day=f'2026-0{part+1}-02'
                writer.writerow(['A',day,journal,'001','001','신한은행',1,0,'신한은행','000-01','V1','D1','P1','작성자'])
                writer.writerow(['A',day,journal,'002','003','매출',0,1,'','','V1','D1','P1','작성자'])
    started=time.perf_counter();value=load_batch(files,profile(),default_catalog(),False);loaded=time.perf_counter()-started
    for row in review_rows(value):apply_review(value,row['key'],{'001':'A1102','003':'R4101'}[row['originalCode']],'확인','명시 시험 설정')
    assert len(value.result['ledger'])==rows and value.result['counts']['errors']==value.result['counts']['pendingRows']==0
    started=time.perf_counter();review=export_session(value,large/'review',exact_folder=True);exported=time.perf_counter()-started
    wb=load_workbook(review/'계정표준화_검토.xlsx',read_only=True,data_only=False)
    try:
        count=0
        for row in wb['원장'].iter_rows(min_row=2,values_only=True):
            assert row[4] in ('001','003');count+=1
        assert count==rows
    finally:wb.close()
    report={'compoundPortableWork':portable,'largeCsv':{'rows':rows,'files':2,'journals':rows//2,'debitTotal':rows//2,'creditTotal':rows//2,
            'pendingRows':0,'conversionErrors':0,'loadSeconds':round(loaded,3),'exportSeconds':round(exported,3),'savedExcelRows':count,
            'streamingReviewExport':True,'leadingZerosPreserved':True},'auditContractLimit':10000}
    (folder/'company-workflow-check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(report,ensure_ascii=False))

if __name__=='__main__':main()

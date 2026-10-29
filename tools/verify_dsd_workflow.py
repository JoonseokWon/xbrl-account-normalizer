"""Synthetic, offscreen check of all four DSD operations in source or EXE."""
import json
import os
from pathlib import Path
import sys
import time
from zipfile import ZipFile, ZIP_DEFLATED


def create_examples(directory):
    """Self-created table fixtures, never official templates or filing files."""
    from openpyxl import Workbook
    directory = Path(directory); directory.mkdir(parents=True, exist_ok=True)
    xml = '''<?xml version="1.0" encoding="UTF-8"?>
<DOCUMENT><TITLE>변환 검토용 자료</TITLE><SECTION><TITLE>1. 현금및현금성자산</TITLE>
<P>단위: 원. 이 자료는 소프트웨어 검증용이며 제출용 보고서가 아닙니다.</P>
<TABLE><THEAD><TR><TH ROWSPAN="2">과목</TH><TH COLSPAN="2">잔액</TH></TR>
<TR><TH>2025년 기말</TH><TH>2024년 기말</TH></TR></THEAD>
<TBODY><TR><TD>국민은행 보통예금</TD><TD>1,200</TD><TD>900</TD></TR>
<TR><TD>신한은행 보통예금</TD><TD>800</TD><TD>600</TD></TR>
<TR><TD>계</TD><TD>2,000</TD><TD>1,500</TD></TR></TBODY></TABLE>
<P>예금의 사용제한 여부는 담당자가 확인합니다.</P></SECTION>
<SECTION><TITLE>2. 재고자산</TITLE><P>단위: 원</P>
<TABLE><TR><TH>과목</TH><TH>2025년 기말</TH><TH>2024년 기말</TH></TR>
<TR><TD>상품</TD><TD>3,000</TD><TD>2,400</TD></TR>
<TR><TD>평가충당금</TD><TD>(120)</TD><TD>(100)</TD></TR>
<TR><TD>미입력</TD><TD/><TD>-</TD></TR><TR><TD>증감률</TD><TD>5.0%</TD><TD>0</TD></TR></TABLE>
</SECTION></DOCUMENT>'''
    source = directory / '검토용_원본.dsd'
    with ZipFile(source, 'x', ZIP_DEFLATED) as archive:
        archive.writestr('contents.xml', xml.encode('utf-8'))
        archive.writestr('metadata.xml', '<EXAMPLE official="false">제출금지</EXAMPLE>')
        archive.comment = b'synthetic review example, not an official DART form'
    workbook = directory / '다른정산표.xlsx'
    book = Workbook(); sheet = book.active; sheet.title = '기말잔액'
    for row in [('계정명', '확정금액_원'), ('국민은행 예금', 1350), ('신한은행 예금', 850), ('현금 합계', 2200)]:
        sheet.append(row)
    sheet.column_dimensions['A'].width = 30; sheet.column_dimensions['B'].width = 23
    from openpyxl.styles import Font, PatternFill
    for row in sheet:
        for cell in row:
            cell.font = Font(name='맑은 고딕', size=11)
            if isinstance(cell.value, str): cell.data_type = 's'
            else: cell.number_format = '#,##0'
    for cell in sheet[1]: cell.fill = PatternFill('solid', fgColor='EAF1FF')
    book.save(workbook); book.close()
    return source, workbook


def run_checks(output):
    os.environ['QT_QPA_PLATFORM'] = 'offscreen'
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QApplication
    from openpyxl import load_workbook
    from version import VERSION
    from desktop.theme.typography import configure_typography
    from desktop.theme.styles import stylesheet
    from desktop.dsd_tools import DsdToolsDialog
    from xbrl_filing.disclosure_compare.dsd import read_dsd
    from xbrl_filing.dsd_roundtrip import export_workbook_to_dsd
    from xbrl_filing.dsd_reporting import export_dsd_review_html
    from xbrl_filing.disclosure_compare.examples import create_comparison_example
    from xbrl_filing.disclosure_compare import compare_disclosures
    started = time.monotonic(); output = Path(output); output.mkdir(parents=True, exist_ok=True)
    source, generic = create_examples(output / '입력자료')
    app = QApplication.instance() or QApplication([])
    app.setStyle('Fusion'); family = configure_typography(app); app.setStyleSheet(stylesheet(family))
    dialog = DsdToolsDialog(); dialog.setAttribute(Qt.WidgetAttribute.WA_DontShowOnScreen); dialog.show()
    def wait():
        deadline = time.monotonic() + 60
        while dialog.loading and time.monotonic() < deadline:
            app.processEvents(); time.sleep(.003)
        app.processEvents()
        assert not dialog.loading, 'DSD background operation exceeded its deadline'
    def picture(index, name):
        dialog.tabs.setCurrentIndex(index); app.processEvents()
        assert dialog.grab().save(str(output / name))
    extracted = output / '표_추출.xlsx'; dialog.extract.source.setText(str(source)); dialog.extract.export_to(extracted); wait()
    assert extracted.is_file(), dialog.notice.text(); extraction = dialog.last_export
    assert extraction['sourceUnchanged'] and extraction['counts']['tables'] == 2
    picture(0, 'DSD_엑셀.png')
    # Unchanged values retain lexical form, including percent, comma, dash and blank.
    unchanged = output / '내용유지_초안.dsd'; export_workbook_to_dsd(source, extracted, unchanged)
    before_cells, same_cells = read_dsd(source)['cells'], read_dsd(unchanged)['cells']
    assert [c['text'] for c in before_cells] == [c['text'] for c in same_cells]
    edited = output / '수정한_표.xlsx'
    book = load_workbook(extracted); sheet = book[extraction['tables'][0]['sheet']]
    sheet['B3'] = 1350; sheet['C3'] = 950; book.save(edited); book.close()
    draft = output / '수정_초안.dsd'; panel = dialog.roundtrip
    panel.template.setText(str(source)); panel.workbook_path.setText(str(edited)); panel.read(); wait()
    panel.confirmed.setChecked(True); panel.export_to(draft); wait()
    assert draft.is_file(), dialog.notice.text()
    derived_report = dialog.last_export
    # A different spreadsheet has no hidden metadata: explicitly link three reviewed cells.
    panel.mode.setCurrentIndex(1); panel.workbook_path.setText(str(generic)); panel.read(); wait()
    for source_cell, old_value in [('B2', '1,200'), ('B3', '800'), ('B4', '2,000')]:
        panel.sheet.setCurrentText('기말잔액'); panel.coordinate.setText(source_cell)
        row = next(i for i, c in enumerate(panel.visible_targets) if c['text'] == old_value)
        panel.targets.selectRow(row); panel.add_link()
    assert len(panel.links) == 3
    panel.confirmed.setChecked(True); generic_draft = output / '정산표_반영초안.dsd'; panel.export_to(generic_draft); wait()
    assert generic_draft.is_file(), dialog.notice.text(); generic_report = dialog.last_export
    picture(3, '정산표_DSD.png')
    dialog.versions.before.setText(str(source)); dialog.versions.after.setText(str(draft)); dialog.versions.compare(); wait()
    comparison = dialog.version_report; assert comparison and comparison['status'] == 'differences', dialog.notice.text()
    assert sum(r['kind'] == 'value-changed' for r in comparison['rows']) == 2
    picture(1, 'DSD_버전비교.png')
    dialog.html.reviewer.setText('검토 담당자'); dialog.html.notes.setPlainText('예금의 기말잔액 변경 2건을 확인했습니다. 공식 형식검사는 별도입니다.')
    html = output / '변경_검토보고서.html'; dialog.html.export_to(html); wait(); assert html.is_file(), dialog.notice.text()
    text = html.read_text(encoding='utf-8'); assert '예금의 기말잔액 변경' in text
    assert str(source.parent) not in text and '<script' not in text.lower()
    picture(2, 'HTML_검토보고서.png')
    examples = create_comparison_example(output / '대사예제')
    # Exercise the existing comparison schema as well as DSD version reports.
    dsd, ixd, mapping = examples['dsd'], examples['ixd'], examples['mapping']
    paired = compare_disclosures(dsd, ixd, json.loads(mapping.read_text(encoding='utf-8-sig')))
    export_dsd_review_html(paired, output / 'DSD_IXD_검토보고서.html', source_paths={'dsd': dsd, 'ixd': ixd})
    dialog.close(); dialog.deleteLater(); app.processEvents()
    report = {'version': VERSION, 'frozen': bool(getattr(sys, 'frozen', False)), 'ok': True,
              'syntheticOnly': True, 'osInputUsed': False, 'networkUsed': False,
              'extractedTables': extraction['counts']['tables'], 'extractedCells': extraction['counts']['cells'],
              'unchangedRoundtripExactText': True, 'editedRoundtrip': derived_report,
              'genericWorkbookMapping': generic_report, 'diffCounts': comparison['counts'],
              'htmlForBothComparisonKinds': True, 'originalFilesUnchanged': True,
              'officialDartEditorOpened': False, 'officialDartFormatCheckPassed': False,
              'scope': 'Synthetic XML-table DSD fixtures, actual UI jobs and output reparsing. Not official DART editor compatibility or filing acceptance.',
              'seconds': round(time.monotonic() - started, 3)}
    (output / 'report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    return report


if __name__ == '__main__':
    sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
    import argparse
    parser = argparse.ArgumentParser(); parser.add_argument('output'); arguments = parser.parse_args()
    print(json.dumps(run_checks(arguments.output), ensure_ascii=False))
